import json
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest


def test_long_ledger_and_token_batches_preserve_last_requirement():
    from desktop.job_brief import build, planning_text
    from desktop.generation import Writer
    from desktop.test_project_quality import ProjectRuntime
    jd='Job Title: Platform Engineer\nRequired Skills\n'+'\n'.join(f'Implement CustomTool{i} reconciliation with retry ownership and error recovery.' for i in range(100))
    brief=build(jd,'','Platform Engineer')
    text=planning_text(brief)
    assert 'CustomTool99' in text
    runtime=ProjectRuntime()
    runtime.token_count=lambda messages:sum(len(m['content'])//3 for m in messages)
    writer=Writer(runtime,jd,lambda *a:None,project_edition=True)
    assert writer.jd==jd
    batches=writer.batches(text.splitlines(),'Plan projects')
    assert len(batches)>1 and '\n'.join(batches)==text
    assert all('source_passage' in r for r in brief['requirements'])


def test_actual_http_retries_share_submission_budget():
    from desktop.runtime import Runtime
    from desktop.generation import Writer
    runtime=Runtime();runtime.token_count=lambda m:100
    requests=[]
    def handler(request):
        requests.append(request)
        return httpx.Response(503,request=request)
    runtime.client=httpx.Client(base_url='http://127.0.0.1',transport=httpx.MockTransport(handler))
    writer=Writer(runtime,'Python',lambda *a:None,project_edition=True)
    writer.calls=26
    with pytest.raises(ValueError,match='budget'):
        runtime.chat([{'role':'user','content':'test'}])
    assert writer.calls==27 and len(requests)==1
    runtime.stop()


def test_ledger_alternatives_and_edit_authority():
    from desktop.job_brief import build
    from desktop.requirement_ledger import report,html
    jd='Required Skills\nUnit testing using pytest, Jest, or equivalent.'
    document={'jd':jd,'job_brief':build(jd,'','Engineer'),'sections':[
        {'key':'skills','text':'pytest, Jest'},
        {'key':'projects','text':'#### Service\n- Tested Python handlers using pytest fixtures with database rollback assertions and malformed request cases before releasing the service.'}]}
    result=report(document)
    assert any(r['status']=='supported' and r['alternatives'] for r in result['rows'])
    assert '#projects-bullet-1' in html(result)
    document['sections'][1]['text']='#### Edited Service\n- Coordinated team meetings.'
    assert not any(r['status']=='supported' for r in report(document)['rows'])


def test_reference_exports_preserve_edited_tables_and_escape(tmp_path):
    from desktop.reference_layout import html,docx
    from docx import Document
    document={'sections':[
        {'key':'header','title':'Header','text':'Test Person\nCloud Engineer'},
        {'key':'skills','title':'Technical Skills','text':'**Cloud:** AWS, EditedTool'},
        {'key':'education','title':'Education','text':'BTech | Example College'},
        {'key':'projects','title':'Relevant Project Experience','text':'### Visible Employer\n#### Test Project\nClient: Confidential client\nTarget Role: Cloud Engineer\n##### Technical Responsibilities\n- Edited <script> text.\n##### Functional Responsibilities\n- Reviewed stakeholder acceptance.'}]}
    page=html(document)
    assert '<table' in page and '&lt;script&gt;' in page and '<script>' not in page
    path=tmp_path/'output.docx';docx(document,path);result=Document(path)
    assert len(result.tables)==3
    text=' '.join(p.text for p in result.paragraphs)+' '.join(c.text for t in result.tables for r in t.rows for c in r.cells)
    assert 'EditedTool' in text and 'Visible Employer' in text and 'Functional Responsibilities' in text
    assert 'C6D9EE' in result._element.xml and 'cantSplit' in result._element.xml


def test_cuda_allocation_failure_tries_partial_then_cpu(monkeypatch,tmp_path):
    import desktop.runtime as module
    (tmp_path/'logs').mkdir();(tmp_path/'runtime-cuda').mkdir();(tmp_path/'runtime-cuda/llama-server.exe').touch()
    (tmp_path/'model.gguf').write_bytes(b'model')
    (tmp_path/'manifest.json').write_text(json.dumps({'files':[]}))
    monkeypatch.setattr(module,'ASSETS',tmp_path);monkeypatch.setattr(module,'DATA',tmp_path)
    monkeypatch.setattr(module,'asset_check',lambda name:tmp_path/'model.gguf' if name.startswith('models') else tmp_path/name)
    monkeypatch.setattr(module,'model_spec',lambda:{'filename':'model.gguf','layers':40})
    monkeypatch.setattr(module.subprocess,'run',lambda *a,**k:SimpleNamespace(returncode=0,stdout='6000'))
    launched=[]
    def launch(args,**kwargs):
        layers=int(args[args.index('-ngl')+1]);launched.append(layers)
        return SimpleNamespace(pid=123,poll=lambda:None if layers==0 else 1,terminate=lambda:None,wait=lambda **k:None)
    monkeypatch.setattr(module.subprocess,'Popen',launch)
    monkeypatch.setattr(module.httpx,'Client',lambda **k:SimpleNamespace(get=lambda *a,**k:SimpleNamespace(is_success=True),close=lambda:None))
    runtime=module.Runtime();monkeypatch.setattr(runtime,'_owns_listener',lambda port:True)
    runtime.start()
    assert launched==[999,20,0] and runtime.model_info['backend']=='llama.cpp-cpu'
    runtime.stop()


def test_example_assistant_does_not_satisfy_separate_vibe_coding_requirement():
    from desktop.requirement_ledger import alternative_groups
    text='AI coding tools such as GitHub Copilot, Cursor, ChatGPT, etc., and Vibe Coding practices'
    groups=alternative_groups(text,['GitHub Copilot','Cursor','ChatGPT','Vibe Coding'])
    assert groups==[['GitHub Copilot','Cursor','ChatGPT']]


def test_parenthetical_framework_preference_keeps_python_mandatory():
    from desktop.job_brief import build
    brief=build('Required Skills\nPython (FastAPI preferred)','','Developer')
    assert any(r['required'] and r['terms']==['python'] for r in brief['requirements'])
    assert any(not r['required'] and 'fastapi' in r['terms'] for r in brief['requirements'])


def test_owned_delivery_is_evidence_but_skills_alone_are_not():
    from desktop.requirement_ledger import report
    doc={'jd':'Required Skills\nGit','sections':[{'key':'projects','text':'#### Service\n- Owned release delivery from Git merge through production rollout, with reviewed changes and documented rollback procedures for application support teams.'}]}
    assert any(r['status']=='supported' for r in report(doc)['rows'])
    doc['sections']=[{'key':'skills','text':'Git'}]
    assert not any(r['status']=='supported' for r in report(doc)['rows'])


def test_rule_patterns_are_readable_in_repair_prompts():
    from desktop.project_rules import facet_display, findings
    assert facet_display(r'(?<!\w)api\ integration(?!\w)')=='api integration'
    assert facet_display(r'\.NET')=='.NET'
    assert any('business records' in w for w in findings('Implemented PostgreSQL handlers before persisting records in encrypted secret stores.'))


def test_summary_rejects_unsubstantiated_spelled_metrics_without_forcing_generic_openers():
    from desktop.project_profile import summary_issues
    projects=[{'tools':'Python','bullets':['Implemented Python services with isolated tests and reviewed release automation.']}]
    warnings=summary_issues(['Engineer with 9+ years of experience building services.','Executed benchmarks raising coverage to ninety-eight percent on fifty thousand records.'],projects,'Engineer',9)
    assert any('quantified' in w for w in warnings)
    assert not any('career-level' in w for w in warnings)
    from desktop.project_profile import polish_summary_opening
    assert polish_summary_opening('Designed and implemented policy controls.')=='Designs and implements policy controls.'


def test_new_template_removes_quarters_preserves_versions_and_old_text():
    from desktop.quality import date_free_document
    doc={'sections':[{'key':'projects','text':'Deployed in Q2 using SQL Server 2022 with 9+ years of experience.'}]}
    assert 'Q2' in date_free_document(doc)['sections'][0]['text']
    doc['template_version']='reference-2'
    text=date_free_document(doc)['sections'][0]['text']
    assert 'Q2' not in text and 'SQL Server 2022' in text and '9+ years' in text


def test_quality_transport_retries_cannot_exceed_two_repair_submissions():
    from desktop.runtime import Runtime
    from desktop.generation import Writer
    runtime=Runtime();runtime.token_count=lambda m:100
    requests=[]
    def handler(request):
        requests.append(request);return httpx.Response(503,request=request)
    runtime.client=httpx.Client(base_url='http://127.0.0.1',transport=httpx.MockTransport(handler))
    writer=Writer(runtime,'Python',lambda *a:None,project_edition=True)
    writer.current_stage='quality_repair'
    with pytest.raises(RuntimeError):runtime.chat([{'role':'user','content':'test'}])
    with pytest.raises(ValueError,match='quality-repair'):runtime.chat([{'role':'user','content':'test'}])
    assert len(requests)==2 and writer.quality_submissions==2
    runtime.stop()


def test_windows_job_retains_four_gib_commit_limit():
    import os,subprocess,sys,ctypes
    if os.name!='nt':pytest.skip('Windows kernel job limits')
    from desktop.windows_job import OwnedJob,Extended
    process=subprocess.Popen([sys.executable,'-c','import time;time.sleep(60)'],creationflags=subprocess.CREATE_NO_WINDOW)
    job=OwnedJob(process)
    try:
        info=Extended();job.kernel.QueryInformationJobObject.argtypes=[ctypes.c_void_p,ctypes.c_int,ctypes.c_void_p,ctypes.c_uint32,ctypes.c_void_p]
        assert job.kernel.QueryInformationJobObject(job.handle,9,ctypes.byref(info),ctypes.sizeof(info),None)
        assert info.JobMemoryLimit==4*1024**3 and info.BasicLimitInformation.LimitFlags&0x2000
    finally:job.close();process.wait(timeout=5)


def test_installer_self_test_uses_isolated_mutex_only_with_separate_data(tmp_path,monkeypatch):
    from desktop import app
    monkeypatch.delenv('RESUME_DESKTOP_SELF_TEST_MS',raising=False)
    monkeypatch.setattr(app,'DATA',tmp_path)
    monkeypatch.setenv('RESUME_DESKTOP_DATA',str(tmp_path))
    assert app.instance_mutex_name()==app.APP_ID
    monkeypatch.setenv('RESUME_DESKTOP_SELF_TEST_MS','300')
    assert app.instance_mutex_name().startswith(app.APP_ID+'.SelfTest.')
    assert app.instance_mutex_name()!=app.APP_ID
