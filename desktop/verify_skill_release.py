"""Offline frozen-worker export and exact skill checklist acceptance harness."""
from __future__ import annotations
import argparse
from collections import Counter
import copy
import csv
from html import unescape
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / 'api')]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--bundle', required=True)
    parser.add_argument('--inputs', nargs='+', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    bundle = Path(args.bundle).resolve()
    out = Path(args.output).resolve(); out.mkdir(parents=True, exist_ok=True)
    from desktop.skill_coverage import present, named_rows
    from desktop.project_profile import is_role_title
    from desktop.requirement_ledger import report as evidence_report
    from desktop.reference_layout import html as preview
    from docx import Document
    paths = [p for directory in args.inputs for p in Path(directory).resolve().glob('*-cold*.json')]
    duplicate_names=Counter(p.stem for p in paths)
    env = os.environ.copy()
    env.update(RESUME_DESKTOP_EDITION='projects', RESUME_DESKTOP_DATA=str(out/'worker-data'),
               HF_HUB_OFFLINE='1', TRANSFORMERS_OFFLINE='1')
    log = (out/'worker.log').open('wb')
    worker = subprocess.Popen([str(bundle/'ResumeOptimizerProjectsWorker.exe'),'--worker'],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=log, text=True, encoding='utf-8',
        env=env, creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
    counter = 0
    def request(kind, payload):
        nonlocal counter
        counter += 1
        worker.stdin.write(json.dumps({'job_id':str(counter),'kind':kind,'payload':payload})+'\n');worker.stdin.flush()
        while True:
            line = worker.stdout.readline()
            if not line: raise RuntimeError('Frozen worker exited during export')
            event = json.loads(line)
            if event['status']=='error': raise RuntimeError(event.get('message',str(event)))
            if event['status']=='complete': return event['result']
    results=[];checklist=[]
    try:
        assets=request('verify',{})
        assert assets['verified']==161, assets
        for path in paths:
            doc=json.loads(path.read_text(encoding='utf-8'))
            key=(path.parent.name+'-'+path.stem) if duplicate_names[path.stem]>1 else path.stem
            sections={s['key']:s['text'] for s in doc['sections']}
            assert doc['skill_coverage_version']=='explicit-skills-1'
            assert doc['model_calls']<=27 and doc['quality_repair_calls']<=2
            assert len([s for s in sections['summary'].splitlines() if s.startswith('- ')])==6
            assert len(doc['projects'])==3 and all(len(p['bullets'])==7 for p in doc['projects'])
            assert len({p['name'].casefold() for p in doc['projects']})==3
            assert all(is_role_title(p['role']) and is_role_title(p['target_role']) for p in doc['projects'])
            assert sections['projects'].count('### Company Name: #Confidential')==3
            assert '\nCompany Name: #Confidential' not in sections['projects']
            assert sections['projects'].count('Role: ')==3 and 'Target Role:' not in sections['projects']
            assert all(p['employer'].casefold() not in (sections['projects']+'\n'+sections['summary']).casefold() for p in doc['projects'])
            # Standalone four-digit workloads are not calendar dates.
            content='\n'.join(sections.values())
            years=[m.group() for m in re.finditer(r'\b(?:19[7-9]\d|20[0-4]\d)\b',content)
                   if not re.match(r'(?i)\s*(?:documents?|records?|claims?|queries|cases?|samples?|tokens?|requests?|tests?|users?|ms|milliseconds?|seconds?)\b',content[m.end():])
                   and not re.search(r'(?i)\b(?:SQL Server|Windows Server|Visual Studio|Office)\s*$',content[:m.start()])]
            assert not years,(key,'calendar years',years)
            assert not re.search(r'\b(?:LPM|LPA|Timezone:)\b','\n'.join(sections.values()))
            assert 'Additional Capabilities' not in sections['skills']
            terms=sorted({t for row in named_rows(doc['job_brief']) for t in row['terms']})
            assert all(present(sections['skills'],t) for t in terms)
            if 'backend-completeness' in key:
                expected=('Python 3','FastAPI','Asynchronous Programming','REST API','Microservices','PostgreSQL','MySQL',
                          'Redis','OAuth2','JWT','Git','Git Rebase','Git Cherry-Pick','Merge Conflict Resolution',
                          'Unit Testing','Integration Testing','pytest','GitHub Copilot','Cursor','Prometheus',
                          'CloudWatch','Logging','CI/CD','Docker','Kubernetes','Kafka','RabbitMQ','Linux','ORM','SQLAlchemy','Vue.js','GraphQL')
                assert all(present(sections['skills'],t) for t in expected),(key,[t for t in expected if not present(sections['skills'],t)])
            review=evidence_report(doc)
            for term in terms:
                rows=[r for r in review['rows'] if term in r['terms']]
                passages=list(dict.fromkeys(p['text'] for r in rows for p in r['evidence'] if term in p['matched_terms']))
                checklist.append({'sample':key,'skill':term,'priority':'mandatory' if any(r['required'] for r in rows) else 'preferred',
                                  'in_skills':True,'implementation_passages':' | '.join(passages) or 'No identifiable implementation passage'})
            rendered={'preview':preview(doc)}
            exports={}
            for kind in ('docx','html'):
                target=out/(key+'.'+kind);started=time.perf_counter()
                result=request('export',{'document':doc,'kind':kind,'destination':str(target)})
                if Path(result['artifact']).resolve()!=target:shutil.copy2(result['artifact'],target)
                exports[kind]={'seconds':round(time.perf_counter()-started,3),'bytes':target.stat().st_size}
                if kind=='docx':
                    exported=Document(target)
                    rendered[kind]=' '.join(c.text for row in exported.tables[0].rows for c in row.cells)
                    metadata=[(row.cells[0].text.strip(),row.cells[1].text.strip())
                              for table in exported.tables for row in table.rows if len(row.cells)>1]
                    assert ('Company Name:','#Confidential') not in metadata,(key,kind,metadata)
                    resume_text=' '.join(p.text for p in exported.paragraphs)+' '.join(c.text for table in exported.tables for row in table.rows for c in row.cells)
                    assert all(p['employer'].casefold() not in resume_text.casefold() for p in doc['projects'])
                    assert sum(label=='Role' for label,_ in metadata)==3,(key,kind,metadata)
                else:rendered[kind]=target.read_text(encoding='utf-8')
            for kind,text in rendered.items():
                if kind!='docx':
                    assert '<td><strong>Company Name:</strong></td><td>#Confidential</td>' not in text,(key,kind)
                    projects=re.search(r'<h2>Relevant Project Experience</h2>(.*?)</section>',text,re.S)
                    assert projects and all(p['employer'].casefold() not in projects[1].casefold() for p in doc['projects'])
                    assert text.count('<td><strong>Role</strong></td><td>')==3,(key,kind)
                    match=re.search(r'<h2>Technical Skills</h2>(.*?)</section>',text,re.S)
                    assert match,(key,kind,'missing skills section')
                    text=unescape(re.sub(r'<[^>]+>',' ',match[1]))
                assert all(present(text,t) for t in terms),(key,kind,[t for t in terms if not present(text,t)])
            edited=copy.deepcopy(doc)
            for section in edited['sections']:
                if section['key']=='skills':section['text']='**Edited Skills:** MANUAL SKILL EDIT'
            for kind in ('docx','html'):
                target=out/(key+'-edited.'+kind)
                result=request('export',{'document':edited,'kind':kind,'destination':str(target)})
                if Path(result['artifact']).resolve()!=target:shutil.copy2(result['artifact'],target)
                if kind=='docx':
                    d=Document(target);text=' '.join(c.text for row in d.tables[0].rows for c in row.cells)
                    assert 'MANUAL SKILL EDIT' in text
                    assert not any(present(text,t) for t in terms)
                else:
                    text=target.read_text(encoding='utf-8')
                    assert 'MANUAL SKILL EDIT' in text and 'Named JD terms missing from Technical Skills' in text
            results.append({'sample':key,'explicit_skills':len(terms),'all_skills_in_preview_docx_html':True,
                'edited_exports_preserved':True,'exports':exports,'coverage':doc['quality']['coverage'],
                'warnings':doc['quality']['warnings'],'profile_warnings':doc['quality']['profile_warnings']})
            (out/'checks.json').write_text(json.dumps({'assets':assets,'samples':results},indent=2),encoding='utf-8')
            print(key,'exports and exact skills verified',flush=True)
        with (out/'skill-evidence-checklist.csv').open('w',newline='',encoding='utf-8-sig') as stream:
            writer=csv.DictWriter(stream,fieldnames=['sample','skill','priority','in_skills','implementation_passages'])
            writer.writeheader();writer.writerows(checklist)
    finally:
        worker.stdin.close();worker.wait(timeout=30);log.close()


if __name__=='__main__':main()
