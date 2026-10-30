"""Side-by-side edition, grouped document, client masking and compatibility checks."""
import copy
import json
import os
import subprocess
import sys
from pathlib import Path
import pytest
from desktop.test_generation import input_resume
from desktop.test_project_quality import ProjectRuntime as FakeRuntime


def make_grouped(input_resume):
    from desktop.project_generation import generate
    from desktop.draft import markdown
    result = generate({"resume":str(input_resume),"jd":"Build Python services using Docker and FastAPI."},FakeRuntime(),lambda *args:None)
    assert result["schema_version"]==6 and result["edition"]=="projects"
    assert result["model_calls"]<=26 and result["repair_calls"]<=3
    assert len(result["roles"])==len(result["projects"])==len(result["experience_blocks"])==3
    assert len({b["role_id"] for b in result["experience_blocks"]})==3
    for role,block,project in zip(result["roles"],result["experience_blocks"],result["projects"]):
        assert role["role_id"]==block["role_id"]==project["role_id"]
        assert block["employer"]==role["employer"] and block["designation"]==project["role"]
        assert project["client"]=="#Confidential"
        assert len(project["bullets"])==7 and len(project["functional"])==2
    text=markdown(result["sections"])
    assert text.count("## Relevant Project Experience\n")==1
    assert "## Professional Experience" not in text and "## Selected Projects" not in text
    assert all(r["employer"] not in text for r in result["roles"])
    assert text.count("### Company Name: #Confidential")==3
    assert "\nCompany Name: #Confidential" not in text and "Source Employer" not in text
    assert "Example University" in text and "2021" not in text
    return result


def test_grouped_exports_use_edits_and_mask_clients(input_resume,tmp_path,monkeypatch):
    from desktop.draft import export
    import desktop.bootstrap
    monkeypatch.setattr(desktop.bootstrap,"DATA",tmp_path)
    (tmp_path/"artifacts").mkdir()
    doc=make_grouped(input_resume)
    section=next(s for s in doc["sections"] if s["key"]=="projects")
    section["text"]=section["text"].replace("\nRole:","\nCompany Name: Private Example Bank\nRole:",1)
    section["text"]+="\n- Edited implementation delivered for Private Example Bank using Python validation and deployment gates.\nPeriod: Jan 2023 - Dec 2024\n"
    for kind in ("docx","html"):
        artifact=Path(export({"document":doc,"kind":kind,"destination":str(tmp_path/("out."+kind))},lambda *a:None)["artifact"])
        if kind=="docx":
            from docx import Document
            output=Document(artifact);text="\n".join(p.text for p in output.paragraphs)+"\n"+"\n".join(c.text for t in output.tables for r in t.rows for c in r.cells)
        else:text=artifact.read_text(encoding="utf-8")
        assert text.casefold().count("relevant project experience")==1
        assert "Edited implementation delivered" in text and "Private Example Bank" not in text
        assert "2023" not in text and "2024" not in text and "2021" not in text
        assert "North Labs" not in text and "#Confidential" in text
        assert "Company Name:" in text and "Role" in text and "Target Role" not in text


def test_client_mask_keeps_employers_and_technology():
    from desktop.project_document import mask_text
    text="### North Labs\nDesignation: AWS Engineer\n#### Claims Hub\nClient: Private Bank\nTools: AWS, GitHub\n- Delivered for Private Bank using AWS services."
    masked=mask_text(text)
    assert "Private Bank" not in masked and "### North Labs" in masked
    assert "Tools: AWS, GitHub" in masked
    assert "AWS services" in mask_text("Client: AWS\n- Delivered AWS services using audited access controls.")
    edited=mask_text("### North Labs\nCompany name:\nRole: AWS Engineer\nTools: AWS",company_label=True)
    assert 'Company name: #Confidential' in edited and 'Role: AWS Engineer' in edited


def test_new_company_mask_covers_titles_and_summary_but_old_drafts_keep_names(input_resume):
    from desktop.project_document import assemble, sanitize
    doc = make_grouped(input_resume)
    name = doc['roles'][0]['employer']
    doc['projects'][0]['name'] = name + ' Delivery Platform'
    summary = next(s for s in doc['sections'] if s['key'] == 'summary')
    summary['text'] += '\n- Delivered systems at ' + name + '.'
    updated = assemble(doc)
    text = '\n'.join(s['text'] for s in updated['sections'])
    assert name not in text and '### Company Name: #Confidential' in text
    old = copy.deepcopy(doc)
    old['client_label_version'] = 'company-confidential-1'
    old = assemble(old)
    old_text = next(s['text'] for s in old['sections'] if s['key'] == 'projects')
    assert '### ' + name in old_text
    assert next(s['text'] for s in sanitize(old)['sections'] if s['key'] == 'projects') == old_text


def test_restore_removes_only_duplicate_default_company_rows(input_resume):
    from desktop.project_document import sanitize
    doc = make_grouped(input_resume)
    projects = next(s for s in doc['sections'] if s['key'] == 'projects')
    projects['text'] = projects['text'].replace('\nRole:', '\nCompany Name: #Confidential\nRole:', 1)
    projects['text'] = projects['text'].replace('\nRole:', '\nCompany Name: Edited Client\nRole:', 1)
    result = sanitize(doc)
    text = next(s['text'] for s in result['sections'] if s['key'] == 'projects')
    assert text.count('### Company Name: #Confidential') == 3
    assert text.count('\nCompany Name: #Confidential') == 1  # edited client remains a masked row
    assert '\nCompany Name: #Confidential\nCompany Name:' not in text
    assert '\nCompany Name: Edited Client' in projects['text']  # input is untouched


def test_bootstrap_identities_are_isolated():
    root=Path(__file__).resolve().parents[1]
    code="import json; from desktop import bootstrap as b; from desktop import edition as e; print(json.dumps([e.APP_ID,e.APP_LABEL,str(b.DATA),b.worker_command()]))"
    results=[]
    for edition in ("legacy","projects"):
        env=os.environ.copy();env["RESUME_DESKTOP_EDITION"]=edition;env.pop("RESUME_DESKTOP_DATA",None)
        done=subprocess.run([sys.executable,"-c",code],cwd=root,env=env,capture_output=True,text=True,check=True)
        results.append(json.loads(done.stdout))
    assert results[0][0]=="ResumeOptimizerDesktop" and results[1][0]=="ResumeOptimizerProjectEdition"
    assert results[0][2]!=results[1][2]
    assert results[1][3][-2].endswith("project_main.py")


def test_no_automatic_cross_edition_draft_restore(input_resume,tmp_path,monkeypatch):
    from desktop.storage import atomic_json,load_draft
    import desktop.edition as edition
    doc=make_grouped(input_resume)
    path=tmp_path/"draft.json";atomic_json(path,{"document":doc})
    monkeypatch.setattr(edition,"PROJECT_EDITION",False)
    assert load_draft(path) is None
    monkeypatch.setattr(edition,"PROJECT_EDITION",True)
    assert load_draft(path)["document"]["edition"]=="projects"


def test_new_gui_has_one_grouped_editor_and_undo(input_resume,tmp_path,monkeypatch):
    from PySide6.QtWidgets import QApplication
    from desktop.app import Window
    import desktop.app as module
    monkeypatch.setattr(module,"DATA",tmp_path)
    monkeypatch.setattr(module,"load_draft",lambda *a:None)
    monkeypatch.setattr(module,"APP_LABEL","Resume Optimizer Project Edition")
    app=QApplication.instance() or QApplication([])
    window=Window()
    try:
        window.set_document(make_grouped(input_resume))
        titles=[window.sections.item(i).text() for i in range(window.sections.count())]
        assert titles.count("Relevant Project Experience")==1
        assert "Professional Experience" not in titles and "Selected Projects" not in titles
        index=titles.index("Relevant Project Experience");window.sections.setCurrentRow(index)
        old=window.editor.toPlainText();window.editor.insertPlainText("Edited ");window.editor.undo()
        assert window.editor.toPlainText()==old
        window.autosave();assert (tmp_path/"last-draft.json").is_file()
    finally:window.close()


def test_grouped_generation(input_resume):
    make_grouped(input_resume)



def test_employer_names_never_use_client_placeholder(input_resume):
    from desktop.project_generation import generate
    runtime=FakeRuntime();original=runtime.chat
    def response(messages,**kwargs):
        value=json.loads(original(messages,**kwargs))
        if 'roles' in value:
            for role in value['roles']:role['employer']='Confidential client'
        return json.dumps(value)
    runtime.chat=response
    result=generate({'resume':str(input_resume),'jd':'Build Python services.'},runtime,lambda *a:None)
    names=[r['employer'] for r in result['roles']]
    assert len(set(names))==3 and all('confident' not in n.casefold() for n in names)
    assert result['repair_calls']<=3
