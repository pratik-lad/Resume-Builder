"""Desktop regression checks. Assets are optional for pure tests."""
from __future__ import annotations

import copy
import json
import os
import socket
import subprocess
import sys
import time
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "api"))
os.environ["NUM_CTX"] = "4096"
os.environ["CTX_SIZE"] = "8192"
os.environ["EMBED_ENABLED"] = "false"


@pytest.fixture
def resume(tmp_path):
    from docx import Document
    file = tmp_path / "Resume with spaces.docx"
    doc = Document()
    for line in ("Alex Example", "alex@example.test", "Professional Summary",
                 "Python developer building APIs.", "Technical Skills", "Python, FastAPI, Docker",
                 "Professional Experience", "Backend Developer | Example Labs | Jan 2022 - Dec 2024",
                 "- Built Python APIs with FastAPI.", "- Improved throughput by 20%.",
                 "Education", "BSc Computer Science | Example College | 2021"):
        doc.add_paragraph(line)
    doc.save(file)
    return file


def test_fast_source_draft_never_calls_llm(resume, monkeypatch):
    from desktop.draft import generate, markdown
    import desktop.embedding
    from matching import embedder
    from llm import client
    monkeypatch.setattr(client, "chat", lambda *a, **k: pytest.fail("Fast draft called model"))
    monkeypatch.setattr(desktop.embedding, "install", lambda: None)
    monkeypatch.setattr(embedder, "available", lambda: False)
    result = generate({"resume": str(resume), "jd": "Requirements\nBuild Python APIs using FastAPI.\nExperience using Kubernetes.",
                       "mode": "source"}, lambda *args: None)
    text = markdown(result["sections"])
    assert "Example Labs" in text and "Example College" in text
    assert "20%" in text and "Jan 2022 - Dec 2024" not in text
    assert "Kubernetes" not in text
    assert result["source_gap"]["missing_required"] or result["source_gap"]["partial"]


def test_demo_allows_fabrication(resume, monkeypatch):
    from desktop.draft import generate, markdown
    import desktop.embedding
    from matching import embedder
    monkeypatch.setattr(desktop.embedding, "install", lambda: None)
    monkeypatch.setattr(embedder, "available", lambda: False)
    result = generate({"resume": str(resume), "jd": "Requirements\nBuild solutions using Kubernetes and Terraform.",
                       "mode": "demo"}, lambda *args: None)
    text = markdown(result["sections"])
    assert "Delivery Operations Platform" in text
    assert "30%" in text and "terraform" in text.lower()


def test_demo_polish_does_not_apply_source_guard(monkeypatch):
    from desktop.polish import polish
    monkeypatch.setattr("desktop.polish.supported", lambda *args: pytest.fail("Demo was source-restricted"))
    class Runtime:
        def start(self, *args): pass
        def token_count(self, messages): return 100
        def chat(self, messages): return "Invented a Kubernetes platform and saved 75%."
    document = {"sections": [{"key": "summary", "title": "Summary", "text": "Python developer."}],
                "mode": "demo", "jd": "Kubernetes engineer.", "source": "Python developer."}
    result = polish({"document": document, "section": 0}, Runtime(), lambda *args: None)
    assert "75%" in result["sections"][0]["text"]
    assert document["sections"][0]["text"] == "Python developer."


def test_source_polish_rejects_changed_numeric_claims():
    from desktop.polish import supported
    assert not supported("Improved throughput by 20%.", "Improved throughput by 75%.",
                         "Improved throughput by 20% and unrelated tasks improved by 75%.")


def test_export_uses_edited_content_and_updates_report(tmp_path):
    from desktop.draft import export, requirements
    from docx import Document
    document = {"sections": [{"key": "summary", "title": "Summary", "text": "EDITED CONTENT Docker"}],
                "requirements": requirements("Requirements\nExperience with Docker.").model_dump(), "mode": "source"}
    a = export({"document": document, "kind": "docx", "destination": str(tmp_path / "resume.docx")},
               lambda *args: None)
    assert "EDITED CONTENT" in "\n".join(p.text for p in Document(a["artifact"]).paragraphs)
    b = export({"document": document, "kind": "html", "destination": str(tmp_path / "report.html")},
               lambda *args: None)
    report = Path(b["artifact"]).read_text()
    assert "EDITED CONTENT" in report and "Covered" in report


def test_scanned_pdf_is_reported(tmp_path):
    import fitz
    from desktop.draft import generate
    path = tmp_path / "scanned.pdf"
    pdf = fitz.open()
    pdf.new_page()
    pdf.save(path)
    with pytest.raises(ValueError, match="scanned"):
        generate({"resume": str(path), "jd": "Requirements for Python engineer"}, lambda *args: None)


def test_dates_removed_from_both_exports_and_report_tracks_edits(tmp_path):
    from desktop.draft import export
    from docx import Document
    document = {"mode": "demo", "jd": "Build Python APIs using FastAPI and Docker.", "sections": [
        {"key": "summary", "title": "Summary", "text": "Engineer with 9+ years of experience."},
        {"key": "experience", "title": "Experience", "text": "### Engineer | Northstar Labs\nDates: Jan 2023 - Present\n- Built reliable Python services using FastAPI with monitored deployment gates and integration tests for customer workflows."},
        {"key": "projects", "title": "Projects", "text": "### Service Hub\nDuration: 2022 - 2023\nClient: Riverbank Systems"},
        {"key": "education", "title": "Education", "text": "BSc | Example College | 2021"}]}
    outputs = {}
    for kind in ("docx", "html"):
        result = export({"document": document, "kind": kind, "destination": str(tmp_path / ("resume." + kind))}, lambda *a: None)
        path = Path(result["artifact"])
        text = "\n".join(p.text for p in Document(path).paragraphs) if kind == "docx" else path.read_text(encoding="utf-8")
        assert "2023" not in text and "2022" not in text
        assert "2021" in text and "9+ years" in text
        assert "Northstar Labs" in text and "Riverbank Systems" in text
        outputs[kind] = text
    assert "No narrative evidence" in outputs["html"]
    document["sections"][2]["text"] += "\n- Deployed Docker containers using Python and FastAPI with contract tests, image scanning and monitored rollout gates for reliable service delivery."
    result = export({"document": document, "kind": "html", "destination": str(tmp_path / "edited.html")}, lambda *a: None)
    edited = Path(result["artifact"]).read_text(encoding="utf-8")
    assert "Project evidence coverage: 100.0%" in edited


def test_autosave_is_atomic_and_restorable(tmp_path):
    from desktop.storage import atomic_json, load_json
    target = tmp_path / "draft.json"
    atomic_json(target, {"text": "first"})
    atomic_json(target, {"text": "edited"})
    assert load_json(target) == {"text": "edited"}
    assert list(tmp_path.iterdir()) == [target]


def test_missing_asset_and_context_fail_closed(monkeypatch):
    import desktop.runtime as runtime
    from desktop.runtime import Runtime
    monkeypatch.setattr(runtime, "ASSETS", Path("Z:/missing-desktop-assets"))
    with pytest.raises(RuntimeError, match="manifest"):
        runtime.asset_check("missing.gguf")
    model = Runtime()
    monkeypatch.setattr(model, "token_count", lambda messages: 4096)
    with pytest.raises(ValueError, match="too long"):
        model.chat([{"role": "user", "content": "oversized"}])


def test_low_free_ram_does_not_block_model_start_and_runtime_stays_bounded(monkeypatch, tmp_path):
    import desktop.runtime as runtime
    from types import SimpleNamespace
    (tmp_path / "logs").mkdir()
    (tmp_path / "manifest.json").write_text('{"files":[]}')
    monkeypatch.setattr(runtime, "DATA", tmp_path)
    monkeypatch.setattr(runtime, "ASSETS", tmp_path)
    monkeypatch.setattr(runtime, "asset_check", lambda name: tmp_path / name)
    monkeypatch.setattr(runtime.psutil, "virtual_memory", lambda: SimpleNamespace(available=100 * 1024**2))
    class Captured(Exception): pass
    def capture(args, **kwargs):
        assert args[args.index("-b") + 1] == "256"
        assert args[args.index("-ub") + 1] == "64"
        assert args[args.index("--cache-ram") + 1] == "0"
        assert args[args.index("-c") + 1] == "4096"
        assert args[args.index("-np") + 1] == "1"
        raise Captured()
    monkeypatch.setattr(runtime.subprocess, "Popen", capture)
    model = runtime.Runtime()
    try:
        with pytest.raises(Captured): model.start()
    finally:
        model.stop()


def test_selected_model_and_low_memory_runtime_options(monkeypatch, tmp_path):
    import desktop.runtime as runtime
    (tmp_path / "logs").mkdir()
    spec = {"filename": "Different-4B-Q4_K_M.gguf", "sha256": "a" * 64, "disable_repacking": True,
            "sampling": {"temperature": .7, "top_p": .8, "presence_penalty": .5}}
    (tmp_path / "manifest.json").write_text(json.dumps({"spec": {"model": spec}, "files": []}))
    monkeypatch.setattr(runtime, "ASSETS", tmp_path)
    monkeypatch.setattr(runtime, "DATA", tmp_path)
    monkeypatch.setattr(runtime, "asset_check", lambda name: tmp_path / name)
    class Captured(Exception): pass
    def capture(args, **kwargs):
        assert Path(args[args.index("-m") + 1]).name == spec["filename"]
        assert "--no-repack" in args and args[args.index("-ngl") + 1] == "0"
        raise Captured()
    monkeypatch.setattr(runtime.subprocess, "Popen", capture)
    model = runtime.Runtime()
    try:
        with pytest.raises(Captured): model.start()
        assert spec["sha256"] in model.model_identity
        assert model.sampling == spec['sampling']
    finally:
        model.stop()


@pytest.mark.parametrize("filename", ["../outside.gguf", "C:\\model.gguf", "nested/model.gguf", "weights.bin"])
def test_invalid_model_selection_is_rejected(monkeypatch, tmp_path, filename):
    import desktop.runtime as runtime
    (tmp_path / "manifest.json").write_text(json.dumps({"spec": {"model": {"filename": filename}}}))
    monkeypatch.setattr(runtime, "ASSETS", tmp_path)
    with pytest.raises(RuntimeError, match="selection is invalid"):
        runtime.model_spec()


def test_model_specific_sampling_keeps_context_output_and_schema_limits(monkeypatch):
    from desktop.runtime import Runtime
    from types import SimpleNamespace
    model = Runtime()
    monkeypatch.setattr(model, 'token_count', lambda messages: 100)
    model.sampling = {'temperature': .7, 'top_p': .8, 'top_k': 20, 'min_p': 0, 'presence_penalty': .5, 'repeat_penalty': 1.05}
    requests = []
    def post(path, json):
        requests.append(json)
        return SimpleNamespace(raise_for_status=lambda: None,
            json=lambda: {'choices': [{'finish_reason': 'stop', 'message': {'content': '{"description":"complete"}'}}]})
    model.client = SimpleNamespace(post=post)
    schema = {'type':'object', 'properties':{'description':{'type':'string','maxLength':600}}}
    assert model.chat([{'role':'user','content':'overview'}], json_output=schema, seed=42)
    assert all(requests[0][k] == v for k, v in model.sampling.items())
    assert requests[0]['max_tokens'] == 768 and requests[0]['seed'] == 42
    assert requests[0]['response_format']['schema'] == schema


def test_controller_queue_and_cancellation(monkeypatch):
    os.environ["QT_QPA_PLATFORM"] = "offscreen"
    from PySide6.QtWidgets import QApplication
    from desktop.controller import JobController
    app = QApplication.instance() or QApplication([])
    controller = JobController()
    events = []
    controller.event.connect(events.append)
    monkeypatch.setattr(controller, "_start", lambda job: setattr(controller, "active", job))
    assert controller.submit("draft", {})
    assert controller.submit("draft", {})
    assert controller.submit("draft", {}) is None
    assert events[-1]["status"] == "error"
    start = time.perf_counter()
    controller.cancel()
    assert time.perf_counter() - start < 1
    assert controller.active is None and controller.waiting is None
    controller.close()


def test_windows_owned_job_cleans_children():
    if os.name != "nt":
        pytest.skip("Windows kernel ownership")
    from desktop.windows_job import OwnedJob
    process = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(60)"],
                               creationflags=subprocess.CREATE_NO_WINDOW)
    job = OwnedJob(process)
    job.close()
    process.wait(timeout=5)
    assert process.poll() is not None


def test_corrupt_asset_is_rejected(tmp_path, monkeypatch):
    import desktop.runtime as runtime
    monkeypatch.setattr(runtime, "ASSETS", tmp_path)
    (tmp_path / "model.gguf").write_bytes(b"corrupt")
    (tmp_path / "manifest.json").write_text(json.dumps({"files": [
        {"path": "model.gguf", "size": 7, "sha256": "0" * 64}]}))
    with pytest.raises(RuntimeError, match="Corrupt"):
        runtime.asset_check("model.gguf")


def test_idle_model_is_unloaded(monkeypatch):
    from desktop.runtime import Runtime
    runtime = Runtime()
    runtime.process = object()
    runtime.last_used = time.monotonic() - 301
    stopped = []
    monkeypatch.setattr(runtime, "stop", lambda: stopped.append(True))
    runtime.idle()
    assert stopped


def test_listener_ownership_does_not_accept_another_server():
    from desktop.runtime import Runtime
    from types import SimpleNamespace
    model = Runtime()
    model.process = SimpleNamespace(pid=os.getpid())
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        listener.listen()
        assert model._owns_listener(listener.getsockname()[1])
        assert not model._owns_listener(1)


def test_editor_preview_undo_autosave_restart(tmp_path, monkeypatch):
    os.environ["QT_QPA_PLATFORM"] = "offscreen"
    from PySide6.QtWidgets import QApplication
    from desktop import app as desktop_app
    app = QApplication.instance() or QApplication([])
    monkeypatch.setattr(desktop_app, "DATA", tmp_path)
    window = desktop_app.Window()
    document = {"sections": [{"key": "summary", "title": "Summary", "text": "Original Python developer."}],
                "jd": "Requirements\nExperience with Docker.", "requirements": {}, "mode": "demo"}
    window.set_document(document)
    start = time.perf_counter()
    window.editor.moveCursor(__import__("PySide6.QtGui", fromlist=["QTextCursor"]).QTextCursor.MoveOperation.End)
    window.editor.insertPlainText(" EDITED Docker")
    assert time.perf_counter() - start < .1
    window.editor.undo()
    assert "EDITED" not in document["sections"][0]["text"]
    window.editor.redo()
    window.autosave()
    window.refresh_preview()
    assert "EDITED" in window.preview.toPlainText()
    window.close()
    restored = desktop_app.Window()
    assert "EDITED" in restored.document["sections"][0]["text"]
    restored.close()


def test_edit_changes_the_correct_section_and_is_undoable(tmp_path, monkeypatch):
    os.environ["QT_QPA_PLATFORM"] = "offscreen"
    from PySide6.QtWidgets import QApplication
    from desktop import app as desktop_app
    app = QApplication.instance() or QApplication([])
    monkeypatch.setattr(desktop_app, "DATA", tmp_path)
    window = desktop_app.Window()
    document = {"sections": [
        {"key": "summary", "title": "Summary", "text": "Original summary."},
        {"key": "skills", "title": "Skills", "text": "Python"}],
        "jd": "Requirements for Python engineer", "requirements": {}, "mode": "demo"}
    window.set_document(document)
    window.sections.setCurrentRow(1)
    window.sections.setCurrentRow(0)
    window.editor.selectAll()
    window.editor.insertPlainText("Edited fictional summary.")
    assert window.sections.currentRow() == 0
    assert document["sections"][1]["text"] == "Python"
    window.sections.setCurrentRow(1)
    window.sections.setCurrentRow(0)
    window.editor.undo()
    assert document["sections"][0]["text"] == "Original summary."
    window.close()


def test_build_download_resumes_and_verifies_hash(tmp_path):
    import hashlib
    import threading
    from http.server import BaseHTTPRequestHandler, HTTPServer
    from desktop.prepare_assets import download_verified
    content = b"verified model bytes for the build download test"
    target = tmp_path / "model.gguf"
    target.with_name("model.gguf.part").write_bytes(content[:8])
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            offset = int(self.headers.get("Range", "bytes=0-").split("=")[1].split("-")[0])
            self.send_response(206)
            self.send_header("Content-Range", f"bytes {offset}-{len(content)-1}/{len(content)}")
            self.end_headers()
            self.wfile.write(content[offset:])
        def log_message(self, *args):
            pass
    server = HTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        url = f"http://127.0.0.1:{server.server_port}/model"
        download_verified(url, target, hashlib.sha256(content).hexdigest(), len(content), lambda *args: None)
        assert target.read_bytes() == content
        # A verified cached asset is reused even with an unreachable URL.
        download_verified("http://127.0.0.1:1/unreachable", target, hashlib.sha256(content).hexdigest(),
                          len(content), lambda *args: None)
    finally:
        server.shutdown()
        server.server_close()


def test_build_preserves_identical_runtime_files(tmp_path, monkeypatch):
    from desktop.prepare_assets import write_if_changed
    file = tmp_path / "runtime.dll"
    file.write_bytes(b"verified runtime")
    def fail_write(*args, **kwargs):
        pytest.fail("An identical runtime DLL was rewritten")
    monkeypatch.setattr(Path, "write_bytes", fail_write)
    write_if_changed(file, b"verified runtime")
