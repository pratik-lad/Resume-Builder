"""Complete generation contract, offline GUI and reference-output regressions."""
import json
import os
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "api"))
os.environ.update(NUM_CTX="4096", CTX_SIZE="8192", EMBED_ENABLED="false", QT_QPA_PLATFORM="offscreen")


class FakeRuntime:
    def __init__(self):
        self.calls = 0
        self.prompts = []

    def start(self, *args):
        pass

    def token_count(self, messages):
        return len(str(messages).split())

    def chat(self, messages, json_output=False, seed=42):
        assert json_output
        self.calls += 1
        task = json.loads(messages[1]["content"].split("\nPrevious response")[0])["task"]
        self.prompts.append(task)
        if "exactly 3 fictional roles" in task:
            return json.dumps({"headline": "AI Engineer", "roles": [{"title": t, "employer": e} for t,e in
                [("Lead AI Engineer", "North Labs"), ("AI Engineer", "West Labs"), ("Developer", "South Labs")]]})
        if '"summary":string' in task:
            return json.dumps({"summary": "AI engineer delivering secure Python services and RAG platforms with measurable business outcomes. " * 9})
        if '"categories"' in task:
            return json.dumps({"categories": [{"name": f"Category {i}", "items": [f"Tool {i}A", f"Tool {i}B"]} for i in range(6)]})
        if "Project number:" in task:
            index = task.split("Project number: ")[1]
            return json.dumps({"name": f"Claims Intelligence {index}", "description": "A production multi-agent platform combining retrieval, authenticated tools and reviewer workflows to resolve insurance cases with observable decisions and repeatable evaluation. " * 3,
                "role": "AI Engineer", "client": "North Shore Insurance", "platform": "Azure AKS / Linux",
                "tools": "Python, LangGraph, Azure AI Foundry, PostgreSQL", "team_size": "8 people",
                **{key:key + " implemented secure Python processing with authenticated interfaces, replayable evaluation and monitored deployment gates." for key in ("problem","architecture","data_flow","evaluation","deployment","security","outcome")}})
        if "ONLY seven string fields" in task:
            return json.dumps({key:key + " implemented secure Python processing with authenticated interfaces, replayable dataset evaluation and monitored deployment gates." for key in ("problem","architecture","data_flow","evaluation","deployment","security","outcome")})
        if "Write a 40-60 word project description" in task:
            return json.dumps({"description": "A production multi-agent platform combining retrieval, authenticated tools and reviewer workflows to resolve insurance cases with observable decisions and repeatable evaluation. " * 2})
        if '"certifications"' in task:
            from desktop.generation import CERTIFICATIONS
            return json.dumps({"certifications": [CERTIFICATIONS[i] for i in (0,3,6)]})
        if "Return ONLY these four string fields" in task:
            templates = {
                "architecture":"Architected Python services using isolated ingestion workers and asynchronous task queues to route insurance documents reliably through reviewer workflows with auditable event contracts.",
                "ingestion":"Built document parsers using schema validation, duplicate detection and quarantine queues to extract consistent claim records from scanned correspondence and structured policy submissions.",
                "retrieval":"Developed hybrid search using chunked evidence records and query relevance filters to retrieve adjudication guidelines while recording citation provenance for independent reviewer inspection.",
                "security":"Configured workload identities using scoped service credentials and strict tool permissions, rotating encryption keys separately from encrypted document storage to protect confidential customer records.",
                "testing":"Implemented contract tests using malformed payload fixtures and simulated downstream failures to verify idempotent retry behavior, error isolation and stable document routing across releases.",
                "delivery":"Automated release pipelines with image scanning, staged promotion and versioned deployment manifests to support controlled rollback when production health checks detect application regressions.",
                "evaluation":"Evaluated retrieval precision using labeled policy queries and held-out adjudication cases, comparing baseline relevance against candidate settings before approving each revised search configuration.",
                "operations":"Monitored request latency and processing costs using component telemetry, actionable alerts and capacity forecasts to reduce expensive retries and maintain documented operational service targets."
            }
            return json.dumps({key:value + f" Component {self.calls}." for key,value in templates.items() if key in json_output['properties']})
        if "Return ONLY three string fields" in task:
            return json.dumps({
                "requirements":f"Facilitated stakeholder workshops with claims managers to agree business acceptance criteria, exception routing and reviewer responsibilities before approving the delivery backlog for client {self.calls}.",
                "governance":f"Led policy risk reviews with compliance owners to document decision boundaries, escalation thresholds and business approvals for disputed cases before operational rollout for client {self.calls}.",
                "handover":f"Delivered operator training with support runbooks, incident exercises and escalation ownership, securing service handover sign-off and a documented post-launch review schedule for client {self.calls}."})
        if "One bullet per assigned skill" in task:
            import re
            skills = json.loads(re.search(r"including its EXACT name: (\[.*?\])", task).group(1))
            return json.dumps({"bullets":[f"Implemented {skill} integration through authenticated APIs and validated data contracts, using replayable test fixtures and monitored deployment gates to improve operational throughput by 25 percent." for skill in skills]})
        count = 5 if "exactly 5" in task else (3 if "exactly 3 substantive" in task else 4)
        return json.dumps({"bullets": [f"Implemented secure Python integration {self.calls}-{i} with authenticated tools, measured evaluation baselines, automated regression checks and production telemetry to improve reliable stakeholder delivery by 25 percent." for i in range(count)]})


@pytest.fixture
def input_resume(tmp_path):
    from docx import Document
    path = tmp_path / "Input resume with spaces.docx"
    doc = Document()
    for text in ["Alex Example", "Developer", "alex@example.test", "Summary", "Legacy summary must not be copied.",
                 "Experience", "Source Employer", "Education", "BSc Computer Science | Example University | 2021"]:
        doc.add_paragraph(text)
    doc.save(path)
    return path


def test_complete_generation_and_exports(input_resume, tmp_path):
    from desktop.generation import generate
    from desktop.draft import markdown, export
    from desktop.bootstrap import DATA
    (DATA / "artifacts").mkdir(parents=True, exist_ok=True)
    runtime = FakeRuntime()
    result = generate({"resume": str(input_resume), "jd": "Build Python and Azure RAG platforms."}, runtime, lambda *args: None)
    text = markdown(result["sections"])
    assert result["mode"] == "demo" and result["generation"] == "complete-ai"
    assert runtime.calls >= 15
    assert len(result["roles"]) == 3 and len(result["projects"]) == 2
    assert all(len(p["technical"]) >= 8 and len(p["functional"]) == 3 for p in result["projects"])
    assert text.startswith("Alex Example\nAI Engineer\nalex@example.test")
    assert "Source Employer" not in text and "Legacy summary" not in text
    assert "Example University" in text
    assert "Delivery Operations Platform" not in text
    assert all("dates" not in r for r in result["roles"])
    assert "Dates:" not in text and "Jan 2023" not in text
    assert result["evidence"]["required_skill_count"] > 0
    assert not any("One bullet per assigned skill" in task for task in runtime.prompts)
    assert result["model_calls"] <= 24
    certs = next(s for s in result["sections"] if s["key"] == "certifications")["text"]
    assert "Implemented" not in certs
    result["sections"][1]["text"] += " EDITED Docker coverage."
    from docx import Document
    docx = export({"document": result, "kind": "docx", "destination": str(tmp_path / "Result.docx")}, lambda *args: None)
    paragraphs = Document(docx["artifact"]).paragraphs
    assert any("EDITED Docker" in p.text for p in paragraphs)
    assert not any(p.text.strip().isdigit() for p in paragraphs)
    html = export({"document": result, "kind": "html", "destination": str(tmp_path / "Report.html")}, lambda *args: None)
    assert "EDITED Docker" in Path(html["artifact"]).read_text(encoding="utf-8")


def test_bounded_repairs():
    from desktop.generation import Writer
    runtime = FakeRuntime()
    seeds = []
    def invalid(*args, **kwargs):
        seeds.append(kwargs["seed"])
        return "invalid json"
    runtime.chat = invalid
    writer = Writer(runtime, "Python developer", lambda *args: None)
    with pytest.raises(ValueError, match="Could not complete"):
        writer.ask("task", lambda v: v, "summary")
    assert writer.calls == 2
    assert seeds == [42, 43]


def test_dates_removed_without_changing_education_or_experience_length():
    from desktop.quality import date_free_document
    document = {"sections": [
        {"key":"summary","title":"Summary","text":"9+ years of experience."},
        {"key":"experience","title":"Experience","text":"### Engineer | Acme\nDates: Jan 2023 - Present\n- Deployed Python services in 2023 with regression tests and monitored rollout gates improving throughput by 25%."},
        {"key":"projects","title":"Projects","text":"Period: 2020-2022\nBuilt a pipeline (2023) using Python."},
        {"key":"education","title":"Education","text":"BSc | University | 2021"}],
        "roles":[{"title":"Engineer","employer":"Acme","dates":"2023-Present","bullets":[]}],
        "projects":[]}
    cleaned = date_free_document(document)
    assert "2023" not in cleaned["sections"][1]["text"]
    assert "2020" not in cleaned["sections"][2]["text"]
    assert cleaned["sections"][0]["text"] == "9+ years of experience."
    assert cleaned["sections"][3]["text"] == "BSc | University | 2021"
    assert "dates" not in cleaned["roles"][0]
    assert "dates" in document["roles"][0]  # input snapshot is preserved


def test_keyword_list_is_not_project_evidence():
    from desktop.quality import evidence_report
    from desktop.draft import requirements
    doc={"sections":[{"key":"skills","text":"Python, Docker, FastAPI"},
                     {"key":"projects","text":"Tools: Python, Docker, FastAPI, automated deployment, monitoring, engineering, build systems and all other implementation technologies."}]}
    report=evidence_report(doc,requirements("Build Python APIs using FastAPI and Docker."))
    assert report["keyword_coverage_percent"] == 100
    assert report["project_evidence_coverage_percent"] == 0
    doc["sections"][1]["text"] += "\n- Implemented Python APIs using FastAPI with Docker deployment, contract tests and monitored rollout gates to reduce response latency by 25 percent."
    assert evidence_report(doc,requirements("Build Python APIs using FastAPI and Docker."))["project_evidence_coverage_percent"] == 100


def test_evidence_skill_punctuation_equivalence():
    from desktop.quality import contains_skill
    assert contains_skill("Implemented AI Security & Governance controls", "ai security governance")
    assert contains_skill("Implemented AI Security and Governance controls", "ai security governance")


def test_omitted_skill_names_are_categorized_without_creating_evidence():
    from desktop.quality import complete_skill_categories, evidence_report
    from desktop.draft import requirements
    source = [{"name":"AI Frameworks", "items":["LangGraph"]}]
    completed = complete_skill_categories(source, ["langchain"], "Develop using LangChain.")
    assert source[0]["items"] == ["LangGraph"]
    assert any("LangChain" in c["items"] for c in completed)
    report = evidence_report({"sections":[{"key":"skills", "text":"LangChain, LangGraph"}]}, requirements("Develop using LangChain."))
    assert report["keyword_coverage_percent"] == 100
    assert report["evidence_coverage_percent"] == 0


def test_date_removal_preserves_numeric_metrics_and_product_versions():
    from desktop.quality import without_dates
    text = without_dates("Validated 2000 queries using SQL Server 2022 with latency below 2050 ms.\nDates: Jan 2023 - Present")
    assert "2000 queries" in text and "2050 ms" in text and "SQL Server 2022" in text
    assert "2023" not in text
    assert "2023" not in without_dates("Released on 2023-02-15 and on 15 October 2023.")
    assert "-02-15" not in without_dates("Released on 2023-02-15.")


def test_conflicting_project_baselines_are_rejected():
    from desktop.generation import check_design_metrics
    with pytest.raises(ValueError, match="baselines conflict"):
        check_design_metrics({"evaluation":"Measured accuracy against a 45% baseline.",
                              "outcome":"Improved accuracy from a 90% baseline."})
    check_design_metrics({"evaluation":"Measured precision against a 70% baseline.",
                          "outcome":"Improved precision to 85% against a 70% baseline."})


def test_known_storage_and_api_host_errors_are_rejected():
    from desktop.generation import check_technical_consistency, similar_bullet
    with pytest.raises(ValueError, match="Key Vault cannot"):
        check_technical_consistency("Secured sensitive PII and financial data within Azure Key Vault using Entra ID.")
    with pytest.raises(ValueError, match="API handlers"):
        check_technical_consistency("Designed a custom Python API integration layer using Azure DevOps Pipelines to consume embeddings.")
    check_technical_consistency("Stored encrypted documents in Blob Storage with encryption keys stored in Azure Key Vault.")
    assert similar_bullet("Deployed the solution on AKS using Docker images with monitored rollout gates.",
                          "Deployed the verification solution on AKS using Docker images with monitored rollout gates.")


def test_verbose_bullets_are_shortened_only_at_sentence_boundaries():
    from desktop.generation import responsibilities
    first = "Implemented document ingestion using validated schemas and retry queues to route rejected customer submissions into an audited reviewer workflow."
    extra = "Documented every component and reviewed operational ownership with support engineers before service rollout."
    result = responsibilities({"bullets":[first + " " + extra * 5]}, 1)[0]
    assert result.startswith(first) and len(result.split()) <= 55 and result.endswith('.')


def test_failed_evidence_repair_preserves_completed_resume(input_resume):
    from desktop.generation import generate
    runtime = FakeRuntime()
    original = runtime.chat
    def missing_skill(messages, **kwargs):
        task = json.loads(messages[1]["content"].split("\nPrevious response")[0])["task"]
        if "One bullet per assigned skill" in task:
            return json.dumps({"bullets": ["Implemented unrelated service components using documented interfaces with integration tests and monitored rollout gates to improve reliable customer delivery and incident handling."]})
        return original(messages, **kwargs)
    runtime.chat = missing_skill
    result = generate({"resume": str(input_resume), "jd": "Build Python services using Docker and FastAPI."}, runtime, lambda *a: None)
    assert len(result["projects"]) == 2
    assert result["evidence"]["missing_project_evidence"]


def test_report_excludes_recruiter_metadata_and_matches_years():
    from desktop.draft import requirements
    from pipeline.output_coverage import evaluate
    req = requirements("Gen AI Lead\nExperience-6 to 9 years\nShift-1 - 10 PM / 2 - 11 PM\nBuild Python APIs using FastAPI.")
    assert not any("Shift" in r.text or r.text == "Gen AI Lead" for r in req.requirements)
    coverage = evaluate("Gen AI Lead with 9+ years of experience. Built Python APIs using FastAPI.", req.requirements)
    assert coverage.required_percent == 100


def test_cross_section_duplicates_are_repaired(input_resume):
    from desktop.generation import generate
    import re
    runtime = FakeRuntime()
    original = runtime.chat
    def repeated(messages, **kwargs):
        task = json.loads(messages[1]["content"].split("\nPrevious response")[0])["task"]
        if "NEW distinct responsibilities" in task:
            count = int(re.search(r"exactly (\d+)", task).group(1))
            runtime.calls += 1
            return json.dumps({"bullets": [f"Rebuilt deployment {runtime.calls}-{i} with isolated tenant credentials, asynchronous integration workers, reproducible quality evaluation and audited production approval gates to reduce processing cost by 18 percent." for i in range(count)]})
        output = original(messages, **kwargs)
        if "exactly 5 detailed responsibilities" in task:
            value = json.loads(output)
            value["bullets"] = [re.sub(r"integration \d+-", "integration same-", b) for b in value["bullets"]]
            return json.dumps(value)
        return output
    runtime.chat = repeated
    result = generate({"resume": str(input_resume), "jd": "Python engineer"}, runtime, lambda *args: None)
    bullets = [b for r in result["roles"] for b in r["bullets"]]
    assert len(bullets) == 15 and len(set(bullets)) == 15
    text = result["sections"][3]["text"]
    assert "Rebuilt deployment" in text


def test_clean_source_page_numbers_and_tables(tmp_path):
    from desktop.generation import read_source
    from docx import Document
    path = tmp_path / "Table education.docx"
    doc = Document()
    for text in ["Alex Example", "Python Engineer", "alex@example.test", "Education"]:
        doc.add_paragraph(text)
    table = doc.add_table(rows=2, cols=3)
    for row, vals in zip(table.rows, [["Degree", "Institute", "GPA"], ["BSc", "Example University", "8/10"]]):
        for cell, val in zip(row.cells, vals): cell.text = val
    doc.add_paragraph("1")
    doc.save(path)
    _,name,contacts,edu = read_source(path)
    assert name == "Alex Example" and contacts == ["alex@example.test"]
    assert "BSc | Example University | 8/10" in edu
    assert not edu.endswith("\n1")


def test_simple_gui(input_resume, monkeypatch, tmp_path):
    from PySide6.QtWidgets import QApplication, QPushButton
    from desktop.app import Window
    monkeypatch.setattr("desktop.app.load_draft", lambda *args: None)
    monkeypatch.setattr("desktop.app.DATA", tmp_path / "isolated-gui-profile")
    app = QApplication.instance() or QApplication([])
    window = Window()
    try:
        assert window.pages.currentIndex() == 0
        labels = [b.text() for b in window.findChildren(QPushButton)]
        assert "Generate resume" in labels
        assert not any("polish" in x.lower() or "fast draft" in x.lower() for x in labels)
        assert not hasattr(window, "mode") and not hasattr(window, "low")
        runtime = FakeRuntime()
        from desktop.generation import generate
        doc = generate({"resume": str(input_resume), "jd": "Python developer"}, runtime, lambda *args: None)
        window.set_document(doc)
        assert window.pages.currentIndex() == 1 and window.editor.isHidden()
        window.edit_button.setChecked(True)
        assert not window.editor.isHidden()
        window.sections.setCurrentRow(1)
        before = window.editor.toPlainText()
        window.editor.insertPlainText("EDITED ")
        window.editor.undo()
        assert window.editor.toPlainText() == before
        window.refresh_preview()
        assert "Alex Example" in window.preview.toPlainText()
    finally:
        window.close()


def test_design_evaluation_wording_does_not_abort_resume(input_resume):
    from desktop.generation import generate
    runtime = FakeRuntime()
    original = runtime.chat
    def loose_design(messages, **kwargs):
        result = json.loads(original(messages, **kwargs))
        if "ONLY seven string fields" in messages[1]["content"]:
            result["evaluation"] = "Measured response accuracy and processing latency"
            result["problem"] = "Reduce manual processing"
        return json.dumps(result)
    runtime.chat = loose_design
    result = generate({"resume": str(input_resume), "jd": "Build Python services using Docker and FastAPI."}, runtime, lambda *args: None)
    assert result["model_calls"] == 19
    assert all("held-out test cases" in p["evaluation"] for p in result["projects"])
    assert all(p["problem"].endswith(".") for p in result["projects"])


def test_global_validation_retry_budget():
    from desktop.generation import Writer
    runtime = FakeRuntime()
    attempts = {}
    def chat(messages, **kwargs):
        stage = messages[1]["content"].split("Previous response")[0].strip()
        attempts[stage] = attempts.get(stage, 0) + 1
        return '{}' if attempts[stage] == 2 else 'invalid json'
    runtime.chat = chat
    writer = Writer(runtime, "Python engineer", lambda *args: None)
    for index in range(3):
        assert writer.ask("task " + str(index), lambda v: v, "test") == {}
    with pytest.raises(ValueError, match="Could not complete"):
        writer.ask("fourth task", lambda v: v, "test")
    assert writer.calls == 7 and writer.repairs_left == 0


def test_conflicting_design_metrics_preserve_resume_with_review_note():
    from desktop.generation import normalize_design
    from desktop.quality import PROJECT_DETAILS
    value = {key:"Implemented secure processing using validated interfaces and monitored delivery gates." for key in PROJECT_DETAILS}
    value.update(evaluation="Measured accuracy on held-out cases against a 45% baseline.",outcome="Improved accuracy from a 90% baseline.")
    notes=[]
    result=normalize_design(value,notes,"Claims workflow")
    assert result['outcome']==value['outcome'] and len(notes)==1


@pytest.mark.parametrize("bad_categories", [
    [{"name":"Languages", "items":["Python", "", " ", None, "C#", "Go", "R", "C++", "Python"]}],
    [{"name":"", "items":["", None]}, {"name":"Platforms", "items":[]}],
    [], None, [None, {}, {"name":"Tools", "items":None}],
])
def test_empty_skill_entries_never_abort_generation(input_resume, bad_categories):
    from desktop.generation import generate
    runtime=FakeRuntime()
    original=runtime.chat
    def skills(messages, **kwargs):
        if '"categories"' in json.loads(messages[1]['content'].split('\nPrevious response')[0])['task']:
            runtime.calls += 1
            return json.dumps({'categories':bad_categories})
        return original(messages, **kwargs)
    runtime.chat=skills
    result=generate({'resume':str(input_resume),'jd':'Build Python APIs using Docker and FastAPI.'},runtime,lambda *args:None)
    assert result['model_calls']==19 and result['repair_calls']==0
    text=next(s['text'] for s in result['sections'] if s['key']=='skills')
    assert 'Python' in text and 'Docker' in text and 'FastAPI' in text
    assert len(result['projects'])==2


def test_short_skills_and_empty_categories_are_normalized():
    from desktop.generation import normalize_skills
    assert normalize_skills({'categories':[None, {'name':'Languages','items':['', 'C', 'R', 'Go', 'C#','F#','C++',None,'R','N/A']} ]}) == [
        {'name':'Languages','items':['C','R','Go','C#','F#','C++']}]
    assert normalize_skills({'categories':[{'name':'','items':'Python, Docker; FastAPI'}]})==[
        {'name':'Additional Skills','items':['Python','Docker','FastAPI']}]



def test_empty_skills_without_recognized_jd_terms_use_reviewable_fallback(input_resume):
    from desktop.generation import generate
    runtime=FakeRuntime();original=runtime.chat
    def response(messages,**kwargs):
        task=json.loads(messages[1]['content'].split('\nPrevious response')[0])['task']
        if '"categories"' in task:
            runtime.calls+=1
            return '{"categories":[]}'
        return original(messages,**kwargs)
    runtime.chat=response
    result=generate({'resume':str(input_resume),'jd':'We seek a dependable colleague.'},runtime,lambda *args:None)
    assert result['model_calls']==19
    assert any('Review skills:' in note for note in result['warnings'])
    assert 'Professional Skills:' in next(s['text'] for s in result['sections'] if s['key']=='skills')



def test_project_identity_cosmetics_are_repaired_without_model_retry():
    from desktop.generation import normalize_project_identity
    value={'name':'North Cloud','client':'North Cloud','role':'AWS Engineer','platform':'AWS','tools':'Terraform, EKS','team_size':'8-10 engineers'}
    first=normalize_project_identity(value,[],'AWS Cloud Engineer')
    assert first['client']=='North Cloud' and first['name']!='North Cloud' and first['team_size']=='8'
    second=normalize_project_identity(value,[first],'AWS Cloud Engineer')
    assert second['name']!=first['name']
    value.update(name='John D.',team_size='not specified')
    third=normalize_project_identity(value,[],'AWS Cloud Engineer')
    assert third['name']=='AWS Cloud Engineer Workflow Modernization' and third['team_size']=='8'
