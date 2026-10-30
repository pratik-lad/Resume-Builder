"""Cross-domain regressions from two real JDs and edited resume text."""
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'api'))

from desktop.job_brief import build, planning_text
from desktop.project_rules import ensure_mandatory_skills, findings, report

FIXTURES = Path(__file__).parent / 'fixtures'


def brief(name, source):
    return build((FIXTURES / (name + '-jd.txt')).read_text(encoding='utf-8'), source, 'Professional')


def test_conflicting_yoe_and_recruiter_data_are_separate():
    from desktop.project_profile import target_role
    assert target_role((FIXTURES/'cloud-security-jd.txt').read_text(encoding='utf-8'),'Other') == 'Senior Cloud & Infrastructure Security Engineer DevSecOps'
    ai = brief('forward-ai','Engineer with 8+ years of experience.')
    security = brief('cloud-security','Engineer with 6+ years of experience.')
    assert ai['experience']['overall_years'] == 8
    assert security['experience']['overall_years'] == 8
    assert any('conflicting' in w for w in ai['experience']['warnings'])
    for item in (ai, security):
        prompt = planning_text(item)
        assert '1.4 LPM' not in prompt and 'Timezone:' not in prompt
        assert not any(line.startswith(('Month 1:', 'Month 2:', 'Weeks 1')) for line in prompt.splitlines())
        assert any(r['kind']=='logistics' and 'Budget:' in r['source_passage'] for r in item['requirements'])
    assert any(r['required'] and r['text']=='Terraform' for r in security['requirements'])


def test_required_skills_are_placed_and_example_tools_are_alternatives():
    security = brief('cloud-security','Engineer with 6+ years of experience.')
    categories = ensure_mandatory_skills([{'name':'Cloud & Infrastructure','items':['Git','Kubernetes']}],security,[])
    groups={c['name']:c['items'] for c in categories}
    assert 'Git' in groups['Version Control & Collaboration']
    assert all(name in groups['Cloud Platforms'] for name in ('AWS','Azure','GCP'))
    assert 'Terraform' in groups['Infrastructure, Containers & CI/CD']
    assert 'Cilium' in groups['Kubernetes & Policy Enforcement']
    assert 'Istio' in ' '.join(v for c in categories for v in c['items'])
    ai = ensure_mandatory_skills([],brief('forward-ai',''),[])
    items=' '.join(v for c in ai for v in c['items'])
    assert all(term in items for term in ('Python','NumPy','Pandas','SFT','RLHF','RAG'))
    assert 'Next.js' not in items and 'CTO' not in items
    salesforce=ensure_mandatory_skills([],build('Position: Salesforce Developer\nRequired Skills:\nSalesforce','','Salesforce Developer'),
                                     [{'bullets':['Implemented customer self-service using Experience Cloud with authenticated case workflows.']}])
    assert any('Experience Cloud' in c['items'] and c['name']=='Salesforce Architecture & Business Applications' for c in salesforce)
    merged=ensure_mandatory_skills([
        {'name':'Version Control & Delivery','items':['Reusable engineering standards']},
        {'name':'Version Control & Collaboration','items':['Git']},
        {'name':'AI, GenAI & Retrieval','items':['Fine-tuning pipelines']},
        {'name':'AI Data, Models & Evaluation','items':['LLM']}],{'requirements':[]},[])
    assert len(merged)==2
    assert next(c['items'] for c in merged if c['name']=='Version Control & Collaboration')==['Reusable engineering standards','Git']


def test_result_review_requires_method_window_and_attributable_change():
    weak='- Reduced incident rate from 15% to 5% on 1000 test cases.'
    assert any('measurement method' in w and 'measurement window' in w for w in findings(weak))
    specific='- Reduced incident rate from 15% to 5% on 1000 replayed test cases over four weekly runs after implementing deduplicated alert routing.'
    assert not any('Result needs' in w for w in findings(specific))
    assert any('SAST' in w for w in findings('- Used SAST for container image scanning in CI/CD.'))
    assert any('arithmetic' in w for w in findings('- Reduced vulnerabilities by 92% (baseline: 48 per week; final: 5 per week) across 1000 images.'))
    assert not any('arithmetic' in w for w in findings('- Reduced vulnerabilities by 90% (baseline: 48 per week; final: 5 per week) across 1000 images.'))


def test_report_recomputes_skill_gap_after_edit():
    jd='Position: Cloud Engineer\nRequired Skills:\n- Git\n- AWS'
    item=build(jd,'','Cloud Engineer')
    document={'jd':jd,'job_brief':item,'sections':[
        {'key':'skills','text':'**Version Control & Collaboration:** Git\n**Cloud Platforms:** AWS'},
        {'key':'projects','text':'#### Delivery\n- Implemented AWS deployment checks using Git pull requests and reviewed rollback evidence.'}]}
    assert not report(document)['skills_review']
    document['sections'][0]['text']='**Cloud Platforms:** AWS'
    assert 'git' in {term.casefold() for term in report(document)['skills_review']}
    quality_jd='Position: Software Engineer\nRequired Skills:\nCode quality'
    quality_doc={'jd':quality_jd,'job_brief':build(quality_jd,'','Software Engineer'),
                 'sections':[{'key':'skills','text':'**Testing:** Code Quality'}]}
    assert not report(quality_doc)['skills_review']
    assert not any(c['status']=='supported' for c in report(quality_doc)['capabilities'])


def test_fitting_brief_is_counted_once_without_losing_requirement_rows():
    from desktop.generation import Writer
    from desktop.test_project_quality import ProjectRuntime
    runtime=ProjectRuntime();counted=[]
    def count(messages):
        counted.append(messages)
        return 800
    runtime.token_count=count
    writer=Writer(runtime,'Platform Engineer',lambda *args:None,project_edition=True)
    rows=['Implement Tool'+str(i)+' with explicit validation.' for i in range(40)]
    assert writer.batches(rows,'Plan the career')==['\n'.join(rows)]
    assert len(counted)==1
    assert writer.calls==0


def test_education_header_stays_with_first_qualification(tmp_path):
    from desktop.reference_layout import docx
    from docx import Document
    document={'sections':[{'key':'education','title':'Education',
                          'text':'Degree | Institute | GPA\nBSc | Example College | 8/10\nMSc | Example University | 9/10'}]}
    path=tmp_path/'education.docx';docx(document,path)
    result=Document(path)
    assert result.tables[0].rows[0].cells[0].paragraphs[0].paragraph_format.keep_with_next
