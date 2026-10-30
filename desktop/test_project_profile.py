"""Phase 1: fictional generation, specific profiles, honest coverage and compatibility."""
import copy
import json
import re
from pathlib import Path

import pytest
from desktop.test_generation import input_resume
from desktop.test_project_quality import ProjectRuntime
from desktop.project_profile import experience_years, certification_candidates, summary_issues


@pytest.mark.parametrize('source,jd,expected',[
    ('Architect with 5+ years of experience designing applications.','Experience: 10 years',10),
    ('Total experience: 8 years','7+ years of experience',8),
    ('BSc | 2021','Experience: 10 years',10),
    ('','6 to 9 years of experience',6),  # Lower bound; no calendar arithmetic.
    ('','6-9 years of experience',6),
    ('','',9),
])
def test_numeric_experience_policy(source,jd,expected):
    assert experience_years(source,jd)==expected


def test_experience_normalization_does_not_rewrite_technical_durations():
    from desktop.project_profile import consistent_years
    text=consistent_years('Engineer with 8+ years of experience; retained audit logs for 2 years.',5)
    assert '5+ years of experience' in text and 'for 2 years' in text


def test_summary_continuation_grammar():
    from desktop.project_profile import profile_continuation
    assert profile_continuation('Designed and delivered reliable platforms.')=='designing and delivering reliable platforms.'
    assert profile_continuation('Developed Python services.')=='developing Python services.'
    assert profile_continuation('Python and FastAPI development.')=='Python and FastAPI development.'
    assert profile_continuation('Databricks Runtime and Unity Catalog.')=='Databricks Runtime and Unity Catalog.'
    assert profile_continuation('Strong expertise in Databricks and Python.')=='Databricks and Python.'


def test_quality_repairs_preserve_distinctness_and_slot_responsibilities():
    from desktop.project_profile import safe_project_repair
    original=['Extracted source records through validation and transformation into owned tables.',
              'Implemented processing using business rules and schema contracts.',
              'Integrated APIs using identifier mappings and retries.',
              'Secured access using OAuth and scoped permissions.',
              'Evaluated latency using a benchmark workload and baseline comparison.',
              'Deployed reusable bundles with rollback and runbooks.']
    duplicate=list(original);duplicate[4]=original[0]
    assert not safe_project_repair(original,duplicate)
    lost_slot=list(original);lost_slot[4]='Mapped incoming business entities into a canonical schema.'
    assert not safe_project_repair(original,lost_slot)
    improved=list(original);improved[4]='Evaluated throughput on replayed test requests against a documented baseline.'
    assert safe_project_repair(original,improved)


@pytest.mark.parametrize('fixture',['databricks-jd.txt','aws-jd.txt','azure-genai-jd.txt'])
def test_profiles_are_bullets_scoped_and_written_last(input_resume,fixture):
    from desktop.project_generation import generate
    jd=(Path(__file__).parent/'fixtures'/fixture).read_text(encoding='utf-8')
    runtime=ProjectRuntime()
    result=generate({'resume':str(input_resume),'jd':jd},runtime,lambda *a:None)
    sections={s['key']:s['text'] for s in result['sections']}
    bullets=[l[2:] for l in sections['summary'].splitlines() if l.startswith('- ')]
    assert len(bullets)==6 and len(set(bullets))==6
    headline=sections['header'].splitlines()[1]
    assert bullets[0].startswith(f"{headline} with {result['experience_years']}+ years of experience")
    assert '**' in sections['skills'] and 'Asynchronous Programming' in sections['skills']
    assert 'Role: '+headline in sections['projects']
    assert sections['projects'].count('Contribution: ')==3
    assert '2021' not in sections['education'] and 'Example University' in sections['education']
    summary_index=next(i for i,t in enumerate(runtime.prompts) if 'professional-summary bullets' in t)
    assert all(i<summary_index for i,t in enumerate(runtime.prompts) if 'improved implementation bullets' in t)
    assert result['model_calls']<=26 and result['quality_repair_calls']<=2
    assert result['mode']=='demo' and result['generation']=='complete-ai'
    assert len(result['roles'])==3
    for prompt in runtime.prompts:
        if 'professional-summary bullets' in prompt:
            assert '"project":' not in prompt
            assert all(p['tools'] in prompt and p['contribution'] in prompt for p in result['projects'])
            assert all(any(b in prompt for b in p['bullets']) for p in result['projects'])
    if 'aws' in fixture:
        assert not any('Genie configuration' in t for t in runtime.prompts)


@pytest.mark.parametrize('fixture',['databricks-jd.txt','aws-jd.txt','azure-genai-jd.txt'])
def test_fresh_generation_seed_and_cache_identity(input_resume,fixture):
    from desktop.project_generation import generate
    jd=(Path(__file__).parent/'fixtures'/fixture).read_text(encoding='utf-8')
    a=generate({'resume':str(input_resume),'jd':jd},ProjectRuntime(),lambda *a:None)
    b=generate({'resume':str(input_resume),'jd':jd},ProjectRuntime(),lambda *a:None)
    assert a['generation_id']!=b['generation_id'] and a['generation_seed']!=b['generation_seed']
    assert all(len({p['name'] for p in d['projects']})==3 for d in (a,b))


def test_retries_keep_generation_identity():
    from desktop.generation import Writer
    runtime=ProjectRuntime();seeds=[]
    def invalid(*args,**kwargs):
        seeds.append(kwargs['seed']);return 'invalid json'
    runtime.chat=invalid
    writer=Writer(runtime,'Python',lambda *a:None,project_edition=True)
    identity=writer.generation_id
    with pytest.raises(ValueError):writer.ask('test',lambda v:v,'test')
    assert writer.generation_id==identity and seeds==[writer.seed,writer.seed+1]


def test_section_cache_is_private_to_one_generation(tmp_path,monkeypatch):
    from desktop.generation import Writer
    import desktop.bootstrap as bootstrap
    monkeypatch.setattr(bootstrap,'DATA',tmp_path)
    runtime=ProjectRuntime();runtime.client=object();calls=[]
    def chat(*args,**kwargs):
        calls.append(kwargs['seed']);return '{"value":"usable"}'
    runtime.chat=chat
    first=Writer(runtime,'Python',lambda *a:None,project_edition=True)
    first.ask('same task',lambda v:v,'same')
    first.ask('same task',lambda v:v,'same')
    assert len(calls)==1
    second=Writer(runtime,'Python',lambda *a:None,project_edition=True)
    second.ask('same task',lambda v:v,'same')
    assert len(calls)==2 and len(list((tmp_path/'cache/generation').glob('*.json')))==2


def test_section_cache_is_private_to_model_checksum(tmp_path, monkeypatch):
    from desktop.generation import Writer
    import desktop.bootstrap as bootstrap
    monkeypatch.setattr(bootstrap, 'DATA', tmp_path)
    runtime = ProjectRuntime(); runtime.client = object(); runtime.model_identity = '4B:hash-one'
    calls = []
    def chat(*args, **kwargs):
        calls.append(kwargs['seed']); return '{"value":"usable"}'
    runtime.chat = chat
    writer = Writer(runtime, 'Python', lambda *a: None, project_edition=True)
    writer.ask('same task', lambda v: v, 'same')
    writer.ask('same task', lambda v: v, 'same')
    assert len(calls) == 1
    runtime.model_identity = '4B:hash-two'
    writer.ask('same task', lambda v: v, 'same')
    assert len(calls) == 2


def test_weak_summary_retries_are_bounded_and_output_survives():
    from desktop.generation import Writer
    from desktop.project_profile import make_summary
    runtime=ProjectRuntime();calls=[]
    def repeated(*args,**kwargs):
        calls.append(kwargs['seed'])
        return json.dumps({key:'Designing Python services using typed interfaces, validated requests and isolated background workers to support reliable enterprise application delivery and repeatable operational procedures across distributed business systems.' for key in kwargs['json_output']['properties']})
    runtime.chat=repeated
    writer=Writer(runtime,'Python',lambda *a:None,project_edition=True)
    projects=[{'name':'Service','role':'Engineer','tools':'Python','bullets':['Built Python services using validation.']*6}]
    bullets,warnings=make_summary(writer,projects,'Engineer',9)
    assert len(bullets)==6 and warnings and len(calls)==4
    assert writer.repairs_left==1 and writer.calls<=27


def test_concepts_preserved_without_recruiter_fragments():
    from desktop.project_rules import complete_skills
    categories=[{'name':'Programming & Backend','items':['Python','Asynchronous Programming','Dependency Injection','Strong','India','Managed applications']}]
    result=complete_skills(categories,'Python FastAPI APIs',preserve_concepts=True)
    text=json.dumps(result)
    assert 'Asynchronous Programming' in text and 'Dependency Injection' in text
    assert 'Strong' not in text and 'India' not in text and 'Managed' not in text


def test_certificates_relevant_not_forced():
    aws=certification_candidates('AWS EKS Terraform')
    azure=certification_candidates('Azure GenAI applications')
    assert any('AWS Certified' in t for t in aws) and not any('Microsoft' in t for t in aws)
    assert any('Microsoft' in t for t in azure) and not any('AWS' in t for t in azure)
    assert not certification_candidates('Technical writing and stakeholder facilitation')


def test_date_policy_preserves_old_drafts_and_numeric_claims():
    from desktop.project_document import sanitize
    old={'edition':'projects','sections':[{'key':'education','text':'BSc Example University | 2021'},
         {'key':'summary','text':'Architect with 5+ years of experience.'}], 'projects':[]}
    original=copy.deepcopy(old)
    assert '2021' in sanitize(old)['sections'][0]['text'] and old==original
    new={**old,'date_policy':'all-calendars'}
    text=' '.join(s['text'] for s in sanitize(new)['sections'])
    assert '2021' not in text and '5+ years' in text and 'Example University' in text


def test_coverage_separates_mandatory_preferred_and_skills_from_evidence():
    from desktop.project_rules import report, report_html
    doc={'jd':'Mandatory Skills\nBuild services using Python.\nPreferred Skills\nDocker',
         'sections':[{'key':'skills','text':'**Backend:** Python, Docker'},{'key':'projects','text':'#### API service\nTools: Python, Docker'}]}
    result=report(doc)
    assert set(result['coverage'])=={'mandatory','preferred'}
    assert all(c['status']=='missing' for c in result['capabilities'])
    assert result['coverage']['mandatory']['keyword_present']>=1
    assert result['coverage']['preferred']['keyword_present']>=1
    assert 'not an external ATS score' in report_html(result)


def test_summary_warnings_do_not_check_uploaded_source():
    projects=[{'name':'Invented','role':'Engineer','tools':'Python','bullets':['Built invented Python services using schema validation.']*6}]
    bullets=['Engineer with 9+ years of experience building Python services.']*6
    issues=summary_issues(bullets,projects,'Engineer',9)
    assert any('Repeated' in w for w in issues)
    assert not any('source' in w.lower() for w in issues)


def test_summary_checks_recognize_implementation_equivalents():
    projects=[{'name':'Pipeline','tools':'PySpark','bullets':['Built ingestion using PySpark transformations and Delta tables, evaluated held-out workloads and authenticated REST/SDK calls with OAuth permissions.']*6}]
    bullets=['Engineer with 9+ years of experience designing ETL platforms using SDKs, security and testing.']
    issues=summary_issues(bullets,projects,'Engineer',9)
    assert not any('technology needs' in w for w in issues)
    issues=summary_issues(['Engineer with 9+ years of experience using FastAPI for delivery.'],projects,'Engineer',9)
    assert any('fastapi' in w.lower() for w in issues)


def test_summary_technology_check_does_not_treat_prose_as_tools():
    projects = [{'name': 'Service', 'tools': 'PostgreSQL, GitHub Actions, OAuth2.0',
                 'bullets': ['Built authenticated SQL transactions and tested deployment pipelines.'] * 6}]
    bullets = ['Engineer with 9+ years of experience using SQL and OAuth 2.0, achieving reliable business workflows with repeatable CI/CD validation and controlled release procedures across enterprise engineering teams.']
    assert not any('technology needs' in w for w in summary_issues(bullets, projects, 'Engineer', 9))


def test_lakehouse_jd_does_not_force_context_serving(input_resume):
    from desktop.project_generation import generate
    runtime=ProjectRuntime()
    generate({'resume':str(input_resume),'jd':'Databricks Engineer\nBuild PySpark ETL pipelines with Delta Lake.'},runtime,lambda *a:None)
    assert 'Collibra' not in runtime.prompts[0] and 'Genie' not in runtime.prompts[0]


def test_system_guidance_is_scoped_to_owned_project(input_resume):
    from desktop.project_generation import generate
    runtime=ProjectRuntime();original=runtime.chat;systems=[]
    def record(messages,**kwargs):
        task=json.loads(messages[1]['content'].split('\nPrevious response')[0])['task']
        if task.startswith('Write this fictional project.'):systems.append(messages[0]['content'])
        return original(messages,**kwargs)
    runtime.chat=record
    jd=(Path(__file__).parent/'fixtures/databricks-jd.txt').read_text(encoding='utf-8')
    generate({'resume':str(input_resume),'jd':jd},runtime,lambda *a:None)
    assert len(systems)==3
    assert 'ONE application database' in systems[0] and 'ONE application database' not in systems[1]
    assert 'Collibra owns glossary' in systems[1] and 'Collibra owns glossary' not in systems[0]
    assert 'Z-ordering is a data-layout optimization' in systems[2]


def test_edited_exports_keep_bullets_bold_categories_and_target_role(input_resume,tmp_path,monkeypatch):
    from desktop.project_generation import generate
    from desktop.draft import export
    import desktop.bootstrap as bootstrap
    from docx import Document
    monkeypatch.setattr(bootstrap,'DATA',tmp_path)
    (tmp_path/'artifacts').mkdir()
    doc=generate({'resume':str(input_resume),'jd':'Python FastAPI Docker'},ProjectRuntime(),lambda *a:None)
    summary=next(s for s in doc['sections'] if s['key']=='summary')
    summary['text']=summary['text'].replace('Strong expertise in','Edited implementation expertise in')
    for kind in ('docx','html'):
        path=Path(export({'document':doc,'kind':kind,'destination':str(tmp_path/('out.'+kind))},lambda *a:None)['artifact'])
        if kind=='docx':
            output=Document(path);text='\n'.join(p.text for p in output.paragraphs)+'\n'+'\n'.join(': '.join(c.text for c in r.cells) for t in output.tables for r in t.rows)
            assert any('Asynchronous Programming' in row.cells[-1].text and any(r.bold for p in row.cells[0].paragraphs for r in p.runs) for t in output.tables for row in t.rows)
            assert len([p for p in output.paragraphs if p.style.name=='Bullet'])>=24
        else:
            text=path.read_text(encoding='utf-8')
            assert '<strong>' in text and '<ul>' in text and '**' not in text
        assert 'Edited implementation expertise' in text and 'Role' in text and 'Target Role' not in text
        assert '2021' not in text and '#Confidential' in text and 'North Labs' not in text


def test_html_report_escapes_edited_text():
    from desktop.project_document import report_shell
    html=report_shell({'sections':[{'key':'summary','title':'Professional Summary','text':'- <script>alert(1)</script>\n**Backend:** Python'}],'jd':'<img src=x>'})
    assert '<script>' not in html and '&lt;script&gt;' in html
    assert '<strong>Backend:</strong>' in html and '&lt;img src=x&gt;' in html


@pytest.mark.parametrize('source,jd,total,specialty', [
    ('5+ years Experience in architecting cloud systems', '8+ years in Software Development, 5+ years of Full-Stack Development using Python, React.JS & SQL/PostgreSQL',8,5),
    ('5+ years of experience','5+ years of Full-Stack Development, 8+ years in Software Development',8,5),
    ('Overall experience: 12 years','8+ years in software development, 5+ years in full stack development',12,5),
    ('','8-12 years in Software Development; 5 to 7 years of Full-Stack Development',8,5),
    ('','5+ years of full-stack development',9,5),
    ('','',9,None),
    ('','8+ years in Software Development, 10+ years of Full-Stack Development',10,10),
])
def test_overall_and_specialty_experience(source,jd,total,specialty):
    from desktop.project_experience import experience_profile
    profile=experience_profile(source,jd)
    assert profile['overall_years']==total
    assert (profile['specialties'][0]['years'] if profile['specialties'] else None)==specialty
    if specialty and specialty>8 and not source:
        assert profile['warnings']


def test_normalization_preserves_specialty_and_technical_durations():
    from desktop.project_experience import experience_profile,normalize_experience
    p=experience_profile('5+ years of experience','8+ years in Software Development, 5+ years of Full-Stack Development')
    text=normalize_experience('Developer with 5+ years of experience; 4+ years of full-stack development; retained logs for 2 years.',p)
    assert '8+ years of experience' in text and '5+ years of full-stack development' in text
    assert 'for 2 years' in text


def test_experience_report_reads_edits_not_metadata():
    from desktop.project_experience import experience_profile,review_experience
    p=experience_profile('5+ years of experience','8+ years in Software Development, 5+ years of Full-Stack Development')
    doc={'experience_profile':p,'sections':[{'key':'summary','text':'Developer with 8+ years of experience in software development, including 5+ years of full-stack development.'}]}
    assert all(r['status']=='stated' for r in review_experience(doc)['rows'])
    doc['sections'][0]['text']='Developer with 5+ years of experience.'
    rows=review_experience(doc)['rows']
    assert rows[0]['status']=='below requirement' and rows[1]['status']=='missing'
    doc['sections'][0]['text']='Developer with 8+ years of experience and 10+ years of full-stack development.'
    assert review_experience(doc)['warnings']


def test_full_stack_planning_and_generation(input_resume):
    from docx import Document
    from desktop.project_generation import generate
    from desktop.project_rules import report,report_html
    path=Document(input_resume);path.add_paragraph('5+ years of experience');path.save(input_resume)
    jd=(Path(__file__).parent/'fixtures/full-stack-jd.txt').read_text(encoding='utf-8')
    runtime=ProjectRuntime()
    doc=generate({'resume':str(input_resume),'jd':jd},runtime,lambda *a:None)
    assert doc['experience_years']==8 and doc['experience_profile']['specialties'][0]['years']==5
    summary=next(s['text'] for s in doc['sections'] if s['key']=='summary')
    assert summary.startswith('- Senior Full Stack Developer with 8+ years of experience in software development, including 5+ years of full-stack development using Python, React and SQL/PostgreSQL;')
    assert any(c['label']=='React application implementation' for c in doc['roles'][0]['assigned_capabilities'])
    assert any(c['label']=='AI-assisted development workflow' for c in doc['roles'][1]['assigned_capabilities'])
    assert any(c['label']=='AWS production application ownership' for c in doc['roles'][2]['assigned_capabilities'])
    assert 'GitHub Copilot' in doc['roles'][1]['stack'] and 'Cursor' not in doc['roles'][1]['stack']
    assert all(r['status']=='stated' for r in report(doc)['experience']['rows'])
    assert 'Experience requirements' in report_html(report(doc))
    assert doc['model_calls']<=27 and doc['quality_repair_calls']<=2
    assert all('Genie configuration' not in p for p in runtime.prompts)


def test_skills_only_do_not_satisfy_full_stack_evidence():
    from desktop.project_rules import report
    jd=(Path(__file__).parent/'fixtures/full-stack-jd.txt').read_text(encoding='utf-8')
    r=report({'jd':jd,'sections':[{'key':'skills','text':'React, Python, FastAPI, PostgreSQL, pytest, Jest, AWS, GitHub Copilot'}]})
    assert all(c['status']=='missing' for c in r['capabilities'])


def test_full_stack_guidance_does_not_leak_to_other_domains():
    from desktop.project_rules import full_stack_capabilities,is_full_stack
    for fixture in ['databricks-jd.txt','aws-jd.txt','azure-genai-jd.txt']:
        jd=(Path(__file__).parent/'fixtures'/fixture).read_text(encoding='utf-8')
        assert not is_full_stack(jd) and not full_stack_capabilities(jd)


def test_domain_total_and_platform_specialty_are_separate():
    from desktop.project_experience import experience_profile,career_opening
    p=experience_profile('6+ years of experience','5\u20138 years of data engineering experience, including at least 2 years hands-on with Databricks')
    assert p['overall_years']==6
    assert p['specialties'][0]['label']=='Databricks' and p['specialties'][0]['years']==2
    other=experience_profile('','8+ years in software development, 5+ years in full-stack development using Java and Angular')
    assert 'Python' not in career_opening('Full Stack Developer',8,other)


def test_duration_conflicts_are_visible_after_edits():
    from desktop.project_experience import experience_profile,review_experience
    doc={'experience_profile':experience_profile('','8+ years in software development'),
         'sections':[{'key':'summary','text':'Developer with 8+ years of experience.'},
                     {'key':'projects','text':'Engineer with 5+ years of experience.'}]}
    assert review_experience(doc)['rows'][0]['status']=='inconsistent'


def test_full_stack_named_tools_and_preferred_aws_priority():
    from desktop.project_rules import classify,full_stack_capabilities,project_assignments
    jd=(Path(__file__).parent/'fixtures/full-stack-jd.txt').read_text(encoding='utf-8')
    aws=next(c for c in full_stack_capabilities(jd) if c['label']=='AWS production application ownership')
    assert not aws['required']
    assert all(not r['required'] for r in classify(jd) if r['text'].startswith('AWS Operations'))
    roles=[{'project_name':str(i),'problem':'A different workflow','stack':'Cursor' if i==1 else 'Python, React','contribution':'AI-assisted development workflow' if i==1 else 'Application engineering'} for i in range(3)]
    modified=jd.replace('GitHub Copilot, Cursor, ChatGPT, etc.','Cursor')
    planned=project_assignments(modified,roles)
    assert 'Cursor' in planned[1]['stack'] and 'GitHub Copilot' not in planned[1]['stack']


def test_full_stack_known_storage_auth_boundaries():
    from desktop.project_rules import findings
    assert any('KMS-managed' in w for w in findings('Used KMS keys stored in AWS Secrets Manager.'))
    assert not findings('Used Secrets Manager for credentials and a KMS key identifier for encryption.')
    assert any('durable database' in w for w in findings('Sent records into an ECS container for persistent storage.'))
    assert not findings('Attached an EFS volume to the ECS container for persistent storage.')
    assert any('human user' in w for w in findings('Authenticated user sessions via OAuth service principals.'))


def test_full_stack_does_not_force_unrequested_retrieval():
    from desktop.project_rules import project_assignments
    roles=[{'project_name':str(i),'problem':'Workflow','stack':'Python, React, LLM','contribution':'Application integration'} for i in range(3)]
    planned=project_assignments('Full Stack Developer: Python React PostgreSQL LLM integrations',roles)
    assert 'LLM' in planned[1]['stack'] and 'vector retrieval' not in planned[1]['stack']


def test_profile_brief_includes_final_test_evidence():
    from desktop.project_profile import brief,profile_continuation
    p={'name':'Web','role':'Developer','tools':'Python','bullets':['Built a service.']*4+['Evaluated API behavior using pytest mocks and Jest component assertions.','Deployed the app.']}
    assert 'pytest' in brief([p])[0]['evaluation']
    assert profile_continuation('Core implementation expertise with proficiency in Python services.')=='Python services.'


def test_availability_is_explicitly_unassessed():
    from desktop.project_rules import report
    jd=(Path(__file__).parent/'fixtures/full-stack-jd.txt').read_text(encoding='utf-8')
    result=report({'jd':jd,'sections':[]})
    rows=[r['text'] for r in result['unassessed_requirements'] if r['kind']=='logistics']
    assert any('India - Remote' in r for r in rows)
    assert any('11:00 pm IST' in r for r in rows)


def test_valid_summary_after_malformed_response_keeps_quality_warnings():
    from desktop.generation import Writer
    from desktop.project_profile import make_summary
    runtime=ProjectRuntime();calls=[]
    def chat(*args,**kwargs):
        calls.append(1)
        if len(calls)==1:
            return 'invalid JSON'
        return json.dumps({key:'Designing ETL services using typed interfaces, validated requests and isolated background workers to support reliable enterprise application delivery and repeatable operational procedures across distributed business systems.' for key in kwargs['json_output']['properties']})
    runtime.chat=chat
    writer=Writer(runtime,'Python',lambda *a:None,project_edition=True)
    projects=[{'name':'Service','role':'Engineer','tools':'Python','bullets':['Built Python services using validation.']*6}]
    bullets,warnings=make_summary(writer,projects,'Engineer',9)
    assert len(bullets)==6 and warnings and len(calls)<=4
    assert writer.profile_metadata['fallback_used'] and all('ETL' not in b for b in bullets)
    assert writer.calls<=27


def test_summary_subject_clauses_become_grammatical_continuations():
    from desktop.project_profile import profile_continuation
    assert profile_continuation('React components manage complex forms.')=='implementing React components that manage complex forms.'
    assert profile_continuation('SQL/PostgreSQL transactions handle critical updates.')=='implementing SQL/PostgreSQL transactions that handle critical updates.'
    assert profile_continuation('React and Python implementation expertise.')=='React and Python implementation expertise.'


def test_generated_ai_skill_and_scanner_wording_normalized_without_extra_calls():
    from desktop.project_rules import clean_skills,normalize,findings
    c=clean_skills([{'name':'AI','items':['LLM-powered authentication and authorization flows','Authentication for LLM applications']}])
    assert c[0]['items']==['LLM API Integration','Authentication for LLM applications']
    source='Deployed a pipeline that scans for vulnerabilities and linting via GitHub Copilot before promotion.'
    fixed=normalize(source,'Python')
    assert 'developer-reviewed Copilot-assisted changes' in fixed and not findings(fixed)


def test_edited_skill_warnings_are_reported_without_rewriting():
    from desktop.project_rules import report
    doc={'jd':'Python Engineer','sections':[{'key':'skills','text':'LLM-powered authentication and authorization flows'}]}
    assert any('deterministic' in w for w in report(doc)['warnings'])
    assert doc['sections'][0]['text']=='LLM-powered authentication and authorization flows'
