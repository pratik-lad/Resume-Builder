import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).parents[1] / "api"))
from pipeline.render import build_ideal_docx


def test_reference_design_and_tables_preserve_content(tmp_path):
    markdown = """# Example Candidate
Senior DevOps Engineer
## Summary
- Built **cloud infrastructure**.
## Technical Skills
- **Cloud:** Azure, AWS
- **Languages:** Python, Go
## Education
| Degree | Institute | CGPA |
| --- | --- | --- |
| BTech | Example College | 7.9 |
## Relevant Project Experience
### Platform Modernization
**Client:** Enterprise
**Tools:** Kubernetes, Terraform
**Team size:** 10
**Project Description:** Cloud delivery platform
#### Technical
- Automated deployments.
## Certifications
- Azure Fundamentals
"""
    doc = Document(build_ideal_docx(markdown, tmp_path / "resume.docx"))
    assert len(doc.tables) == 1
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    for expected in ("Example Candidate", "Azure, AWS", "Example College", "7.9",
                     "Enterprise", "10", "Cloud delivery platform", "Automated deployments."):
        assert expected in text
    assert doc.styles["Normal"].font.name == "Times New Roman"
    assert doc.sections[0]._sectPr.find(qn("w:pgBorders")) is None
    assert 'PAGE' in doc.sections[0].footer._element.xml
    assert 'CBD8DF' in doc._element.xml
    assert doc.tables[0].rows[0]._tr.find('.//' + qn('w:tblHeader')) is not None


def test_unstructured_education_and_unknown_sections_remain_visible(tmp_path):
    doc = Document(build_ideal_docx(
        "# Candidate\n## Education\nPhD in Computing\n## Research Interests\nDistributed systems",
        tmp_path / "resume.docx"))
    assert "PhD in Computing" in [p.text for p in doc.paragraphs]
    assert "Distributed systems" in [p.text for p in doc.paragraphs]


def test_plain_section_labels_project_titles_and_dates_get_resume_layout(tmp_path):
    markdown = """Candidate Name
Data Engineer
candidate@example.com
EDUCATIONAL QUALIFICATION
BTech Computer Science
PROJECTS
Telemetry Platform
Built a monitoring platform.
## Professional Experience
### Engineer | Example Corp
Dates: 01/2022 - Present
- Delivered data pipelines.
### Quality Check (Perform silently before returning)
1. Is every requirement present?
## Certifications
- Cloud certification
"""
    doc = Document(build_ideal_docx(markdown, tmp_path / "layout.docx"))
    paragraphs = {p.text: p for p in doc.paragraphs}
    assert paragraphs["Candidate Name"].style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraphs["Data Engineer"].style.name == "Headline"
    assert paragraphs["EDUCATIONAL QUALIFICATION"].style.name == "SectionHeading"
    assert paragraphs["Telemetry Platform"].style.name == "Subheading"
    role_table = next(table for table in doc.tables
                      if table.cell(0, 0).text == "Engineer | Example Corp")
    assert role_table.cell(0, 1).text == "01/2022 - Present"
    assert role_table.cell(0, 1).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paragraphs["Delivered data pipelines."].style.name == "Bullet"
    assert "1. Is every requirement present?" not in paragraphs
    assert "Cloud certification" in paragraphs


def test_projects_are_grouped_under_matching_employer_in_docx(tmp_path):
    markdown = """# Asha Rao
## Professional Experience
### Engineer | Alpha Corp
Dates: 2023 - Present
- Led platform work.
### Engineer | Beta Corp
Dates: 2021 - 2023
- Built services.
## Selected Projects
### Beta Platform
Employer: Beta Corp
Project Description: Older platform.
- Delivered reporting.
### Alpha Platform
Employer: Alpha Corp
Project Description: Newer platform.
- Delivered telemetry.
## Education
Engineering degree
"""
    doc = Document(build_ideal_docx(markdown, tmp_path / "grouped.docx"))
    xml = doc.element.body.xml
    assert xml.index("Alpha Corp") < xml.index("Alpha Platform") < xml.index("Beta Corp")
    assert xml.index("Beta Corp") < xml.index("Beta Platform") < xml.index("Engineering degree")
    assert "Employer:" not in xml
    assert "SELECTED PROJECTS" not in xml


def test_technical_skills_are_bold_single_line_categories_without_table(tmp_path):
    markdown = """# Asha Rao
## Technical Skills
- **Analytics & BI:** Power BI, Tableau, Excel
- **SQL & Databases:** SQL, PostgreSQL
- **Programming:** Python
## Education
Engineering degree
"""
    doc = Document(build_ideal_docx(markdown, tmp_path / "skills.docx"))
    assert not doc.tables
    skills = [p for p in doc.paragraphs if p.style.name == "SkillLine"]
    assert len(skills) == 3
    assert [p.text.split(":", 1)[0] for p in skills] == [
        "Analytics & BI", "SQL & Databases", "Programming"]
    assert all(p.runs[0].bold for p in skills)
    assert skills[0].text == "Analytics & BI: Power BI, Tableau, Excel"


def test_saved_run_docx_export_drops_repair_echo(tmp_path):
    markdown = """# Kavya Iyer
## Other Skills and Interests
- Organizer of ChennaiJS monthly meetup.
- Achieved Product Impact Award.
Rebuild and tailor this candidate resume to the complete job description.
**MISSING JD SKILLS:** kafka, spark
Return a complete replacement resume.
"""
    doc = Document(build_ideal_docx(markdown, tmp_path / "clean.docx"))
    content = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "EXTRACURRICULAR ACTIVITIES" in content
    assert "ACHIEVEMENTS" in content
    assert "Organizer of ChennaiJS" in content
    assert "Product Impact Award" in content
    assert "Rebuild and tailor" not in content
    assert "MISSING JD SKILLS" not in content
