import json
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parents[1] / "api"))

import main
from config import settings
from pipeline import ideal_resume, jd, optimize
from pipeline import match as match_stage
from schemas import FactLedger, Requirement, RequirementSet, RunRef
import runstore


def test_all_requirements_reach_generation_without_source_matches(monkeypatch):
    captured = {}

    def generate(jd_text, resume_text):
        captured.update(jd=jd_text, resume=resume_text)
        return "## Skills\nKubernetes\n## Education\nPhD"

    monkeypatch.setattr(ideal_resume, "generate", generate)
    reqs = [Requirement(req_id=f"R{i}", text=f"Skill {i}",
                        type="hard_skill", must_have=True) for i in range(40)]
    result = optimize.optimize([], reqs, [], set(), jd_text="Complete JD tail",
                               resume_text="Entry-level background")
    assert all(r.text in captured["jd"] for r in reqs)
    assert "Complete JD tail" in captured["jd"]
    assert "PhD" in result.markdown
    assert result.ucr is None and result.truth_report is None
    assert result.reverted == 0


def test_unverified_mode_scores_missing_requirement_as_covered():
    req_set = RequirementSet(jd_text="Kubernetes", requirements=[
        Requirement(req_id="R1", text="Kubernetes", type="hard_skill", must_have=True,
                    canonical_skills=["kubernetes"])])
    matches, card, gap = match_stage.run(req_set, [], "Entry-level background",
                                         allow_unverified=True)
    assert matches[0].status == "covered"
    assert matches[0].method == "configured"
    assert card.hard_skill_coverage == 100.0
    assert card.jd_alignment_score > 0
    assert not gap.missing_required


def test_optimize_render_without_match_or_truth_gates(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "runs_dir", tmp_path)
    run_id = "idealtest"
    runstore.save_json(run_id, "facts", FactLedger(raw_text="Original background"))
    reqs = RequirementSet(jd_text="Kubernetes and PhD <required>", requirements=[
        Requirement(req_id="R1", text="Kubernetes and PhD", type="qualification", must_have=True)])
    runstore.save_json(run_id, "requirements", reqs)
    # Old safety-era output must not be reused.
    runstore.save_json(run_id, "rewrites", {"rewrites": [], "ucr": 0})
    calls = []
    def generate(*args):
        calls.append(args)
        return "# Ideal Candidate\n## Experience\n- Led Kubernetes delivery for 8 years\n## Education\nPhD"
    monkeypatch.setattr(ideal_resume, "generate", generate)
    result = main.run_optimize(RunRef(run_id=run_id))
    assert main.run_optimize(RunRef(run_id=run_id)).markdown == result.markdown
    assert len(calls) == 1
    rendered = main.run_render(RunRef(run_id=run_id))
    text = "\n".join(p.text for p in Document(rendered.docx_path).paragraphs)
    text += "\n".join(cell.text for t in Document(rendered.docx_path).tables
                      for row in t.rows for cell in row.cells)
    assert "DEMO / FICTIONAL RESUME" not in text
    assert "Kubernetes" in text and "8 years" in text and "PhD" in text
    assert rendered.score_after is None
    report = (tmp_path / run_id / "report.html").read_text(encoding="utf-8")
    assert "Full JD extraction" in report and "&lt;required&gt;" in report
    exported = main.dataset_export()
    assert exported["count"] == 1
    assert exported["rows"][0]["messages"][1]["content"] == reqs.jd_text
    assert exported["rows"][0]["messages"][2]["content"] == result.markdown


def test_direct_endpoint_returns_extraction_and_resume(monkeypatch):
    from schemas import IdealResumeRequest
    requirements = RequirementSet(jd_text="Python", requirements=[
        Requirement(req_id="R1", text="Python", type="hard_skill", must_have=True)])
    monkeypatch.setattr(jd, "extract", lambda text: requirements)
    monkeypatch.setattr(ideal_resume, "generate", lambda *args: "## Skills\nPython")
    result = main.build_ideal_resume(IdealResumeRequest(jd_text="Python"))
    assert result.requirements == requirements
    assert "Python" in result.markdown
    assert "## Professional Summary" in result.markdown
    assert "## Selected Projects" in result.markdown


def test_long_jd_chunks_preserve_tail():
    text = "\n".join(f"Requirement {i}: " + "x" * 100 for i in range(100))
    chunks = jd._chunks(text)
    assert len(chunks) > 1
    assert "Requirement 99" in chunks[-1]
    assert " ".join(" ".join(chunks).split()) == " ".join(text.split())


def test_ai_jd_without_job_title_gets_inferred_title_and_all_sections_recovered():
    text = """Job description
Artificial Intelligence (AI)
# Education
- Microsoft Azure AI Engineer or equivalent AI certification is preferred
## Scope and required competencies
- Implement human review and approval controls for high-impact AI decisions
## Skills & Competencies
- Prompt engineering, retrieval-augmented generation and AI-agent development skills
"""
    assert jd._title_from_jd(text) == "AI Engineer"
    recovered = jd._source_expectations(text)
    lines = [item[0] for item in recovered]
    assert any("Azure AI Engineer" in line for line in lines)
    assert any("human review" in line for line in lines)
    assert any("Prompt engineering" in line for line in lines)


def test_confidential_pipe_role_is_recovered_as_source_employer():
    from matching.experience import extract_entries
    entries = extract_entries("""PROFESSIONAL EXPERIENCE
Senior AI Engineer | #confidential
Whisper (ASR Platform) — Healthcare speech-to-text platform
• Built and deployed models.
""")
    entry = next(item for item in entries if item.role_title == "Senior AI Engineer")
    assert entry.company == "#Confidential"


def test_date_free_confidential_roles_and_unlabelled_source_projects_are_recovered():
    from pipeline import ideal_resume
    source = """PROFESSIONAL EXPERIENCE
Senior AI Engineer | #confidential Whisper (ASR Platform) — Healthcare speech platform
• Took speech models from experimentation to production with Data Science and DevOps.
• Reduced infrastructure cost by 20x using traffic-based Kubernetes autoscaling.
Summary Generator — LLM-based clinical NLP summarization service
• Designed and deployed Gemini summarization microservices for PHI data.
Payer Knowledge Base (PKB) — Enterprise RAG chatbot for insurance claims
• Built a production RAG pipeline with drift monitoring and incremental ingestion.
Associate Software Engineer | #confidential Xpresso.ai — Enterprise MLOps platform
• Built AWS infrastructure with Terraform for enterprise ML workloads.
EDUCATION
B.Tech
"""
    entries = ideal_resume._source_entries(source)
    assert [(entry.role_title, entry.company) for entry in entries] == [
        ("Senior AI Engineer", "#Confidential"),
        ("Associate Software Engineer", "#Confidential"),
    ]
    projects = ideal_resume._source_projects(source)
    assert [name for name, _ in projects] == ["Summary Generator", "Payer Knowledge Base (PKB)"]
    assert all("#Confidential" in value
               for value in ideal_resume._source_project_employers(source).values())


def test_source_context_is_not_reintroduced_as_a_duplicate_fabricated_role():
    from pipeline import ideal_resume
    source = """PROFESSIONAL EXPERIENCE
Senior AI Engineer | #confidential
Whisper (ASR Platform) — Healthcare speech platform
• Owned ASR delivery and production support.
"""
    entries = ideal_resume._source_entries(source)
    assert entries[0].company == "#Confidential"
    generated = """### Senior AI Engineer | #Confidential
- Owned ASR delivery and production support.
### Senior AI Engineer | Healthcare Speech-to-Text Platform (Whisper) #Confidential
- Duplicate role based on the source employer context.
"""
    final = ideal_resume._compact_experience(
        generated, entries, source, [], "AI Engineer and ASR deployment")
    assert "### Senior AI Engineer | #Confidential" in final
    assert "Healthcare Speech-to-Text Platform (Whisper)" not in final
    assert "Duplicate role based on the source employer context" not in final


def test_workflow_has_no_fit_or_claim_gates():
    workflow = json.loads((Path(__file__).parents[1] / "n8n/workflows/resume-optimize.json").read_text())
    assert not any(node["type"] in {"n8n-nodes-base.if", "n8n-nodes-base.stopAndError"}
                   for node in workflow["nodes"])
    assert workflow["connections"]["4 optimize"]["main"][0][0]["node"] == "5 render"
