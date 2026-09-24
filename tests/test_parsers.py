"""Column-detection tests, built from synthetic PDFs generated at test time.

Previously there were no genuine two-column test fixtures at all -- the
"two-column support" claim in the README was unverified. These tests fix that:
one builds an actual two-column PDF and checks the reading order comes out
sidebar-then-body, not interleaved; the other reproduces the exact bug found in
verification (a short education-table row misread as two page columns).
"""
import sys
from pathlib import Path

import fitz
import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "api"))

from parsers.pdf_parser import parse_pdf  # noqa: E402
from parsers.docx_parser import parse_docx  # noqa: E402

PAGE_W, PAGE_H = 612, 792  # US Letter, points


def _make_pdf(tmp_path: Path, name: str, lines: list[tuple[str, float, float]]) -> str:
    """lines: list of (text, x, y) in PDF points, top-down y."""
    doc = fitz.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    for text, x, y in lines:
        page.insert_text((x, y), text, fontsize=10)
    path = tmp_path / name
    doc.save(str(path))
    doc.close()
    return str(path)


def test_genuine_two_column_resume_reads_sidebar_then_body(tmp_path):
    """A real two-column layout: a narrow left sidebar (Skills) running the
    full page height, and a wider right body (Experience) also running the
    full page height. Correct behaviour: every sidebar line before every
    body line, each internally in top-to-bottom order.
    """
    sidebar = [(f"Skill {i}", 50, 60 + i * 40) for i in range(15)]       # x=50, full height
    body = [(f"Experience bullet {i} FastAPI Python", 340, 60 + i * 40)  # x=340, full height
            for i in range(15)]
    path = _make_pdf(tmp_path, "two_column.pdf", sidebar + body)

    doc = parse_pdf(path)
    texts = [l.text for l in doc.lines]

    last_sidebar_idx = max(i for i, t in enumerate(texts) if t.startswith("Skill"))
    first_body_idx = min(i for i, t in enumerate(texts) if t.startswith("Experience"))
    assert last_sidebar_idx < first_body_idx, (
        "sidebar and body were interleaved instead of read column-by-column: "
        f"{texts}"
    )
    assert [t for t in texts if t.startswith("Skill")] == [s[0] for s in sidebar]
    assert [t for t in texts if t.startswith("Experience")] == [b[0] for b in body]


def test_single_column_resume_stays_in_document_order(tmp_path):
    """A normal single-column resume must never be column-split."""
    lines = [(f"Line {i} of a normal single column resume", 50, 60 + i * 20)
             for i in range(25)]
    path = _make_pdf(tmp_path, "single_column.pdf", lines)

    doc = parse_pdf(path)
    texts = [l.text for l in doc.lines]
    assert texts == [l[0] for l in lines]


def test_education_table_is_not_misread_as_two_columns(tmp_path):
    """Regression test for the exact bug found in verification: a resume that
    is single-column EXCEPT for a small education table where the degree
    sits at one x-position and the year/CGPA sits further right on the same
    rows. Before the fix, this could be misclassified as a two-column page
    and the table rows would be reordered into "all left cells, then all
    right cells" instead of staying in row order.

    We build a page with plenty of ordinary single-column body text (all at
    x=50) plus a short 3-row table where the second cell is at x=340 -- the
    same x-position used for a genuine right column above, but confined to a
    handful of rows near the bottom rather than the full page height.
    """
    body = [(f"Experience bullet {i} about Python and FastAPI", 50, 60 + i * 20)
            for i in range(18)]
    table = [
        ("B.Tech Computer Science", 50, 560), ("2020 - 2024", 340, 560),
        ("XYZ Engineering College", 50, 580), ("CGPA 8.5", 340, 580),
        ("Higher Secondary", 50, 600), ("2018 - 2020", 340, 600),
    ]
    path = _make_pdf(tmp_path, "education_table.pdf", body + table)

    doc = parse_pdf(path)
    texts = [l.text for l in doc.lines]

    # The table rows must stay adjacent and in row order -- NOT get split into
    # "all left-side table cells" followed later by "all right-side cells".
    degree_idx = texts.index("B.Tech Computer Science")
    year_idx = texts.index("2020 - 2024")
    assert year_idx == degree_idx + 1, (
        f"table row was split across a false column boundary: {texts[degree_idx:degree_idx+4]}"
    )


def test_short_two_column_block_below_line_count_floor_is_not_split(tmp_path):
    """A page with too few total lines to reasonably judge column structure
    must fall back to plain top-to-bottom order, even if it superficially has
    two x-clusters.
    """
    lines = [("Name Here", 50, 60), ("Contact info", 50, 80),
             ("Skill A", 50, 100), ("Skill B", 340, 100),
             ("Skill C", 50, 120), ("Skill D", 340, 120)]
    path = _make_pdf(tmp_path, "too_short.pdf", lines)

    doc = parse_pdf(path)
    texts = [l.text for l in doc.lines]
    assert texts == [l[0] for l in lines]


def test_docx_keeps_table_between_surrounding_paragraphs(tmp_path):
    doc = Document()
    doc.add_paragraph("Experience")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "ABC Technologies"
    table.cell(0, 1).text = "Software Developer"
    doc.add_paragraph("Education")
    path = tmp_path / "ordered.docx"
    doc.save(path)
    texts = [line.text for line in parse_docx(str(path)).lines]
    assert texts == ["Experience", "ABC Technologies | Software Developer", "Education"]
