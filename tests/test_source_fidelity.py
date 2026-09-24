import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parents[1] / 'api'))

from docx import Document
from pipeline.render import build_docx
from schemas import Rewrite


def test_source_content_survives_incomplete_extraction(tmp_path):
    source = ('Candidate\nemail@example.com | Portfolio | City\n'
              'Profile Summary\nFirst sentence.\nSecond sentence.\n'
              'Projects\nProject name\n2026\nBuilt an API.\n'
              'Education\nUniversity\nDegree\nCGPA: 6.59\n'
              'Achievements & Certifications\n• Certificate – Issuer\n')
    path = build_docx([], [], source, None, tmp_path / 'resume.docx')
    paragraphs = [p.text for p in Document(path).paragraphs]
    assert 'First sentence. Second sentence.' in paragraphs
    for text in ['email@example.com | Portfolio | City', 'Project name', '2026',
                 'University', 'Degree', 'CGPA: 6.59', 'Certificate – Issuer',
                 'Profile Summary', 'Achievements & Certifications']:
        assert text in paragraphs


def test_only_accepted_matching_rewrite_changes_source(tmp_path):
    source = 'Candidate\nProjects\nBuilt an API.\nUnextracted detail.'
    rewrite = Rewrite(fact_id='PRJ-01', target_req='R01', original='Built an API.',
                      final='Developed an API.', status='accepted')
    path = build_docx([], [rewrite], source, None, tmp_path / 'resume.docx')
    text = '\n'.join(p.text for p in Document(path).paragraphs)
    assert 'Developed an API.' in text
    assert 'Unextracted detail.' in text
