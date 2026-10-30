from docx import Document

from docgen import reference_template
from pipeline.presentation import format_resume


def test_display_changes_preserve_experience_content_and_source_companies(tmp_path):
    resume = """# Candidate
candidate2024@example.com | https://example.com/2024/01
## Professional Summary
Engineer with 8 years of experience.
## Technical Skills
SQL Server 2022, ISO 27001, Python
## Professional Experience
### Engineer | Original Corp
Dates: Jan 2018 - Dec 2020
- Improved throughput by 35%.
### Senior Engineer | Added Corp
Dates: Jan 2021 - Present
- Built services handling 2024 transactions per second.
## Selected Projects
### Platform
Employer: Added Corp
Period: 2023 - Present
Duration: 12 months
- Delivered APIs with Python.
## Education
- B.Tech | University | 2017
## Certifications
- Cloud certification, 2024
"""
    result = format_resume(resume, "Candidate\nEngineer at Original Corp")
    assert "### Engineer | Original Corp" in result
    assert "Original Corp #Confidential" not in result
    assert "Added Corp #Confidential" in result
    assert "Dates:" not in result and "Period:" not in result and "Duration:" not in result
    for preserved in ("8 years of experience", "35%", "2024 transactions", "SQL Server 2022",
                      "ISO 27001", "Delivered APIs with Python", "candidate2024@example.com",
                      "https://example.com/2024/01", "B.Tech | University", "Cloud certification"):
        assert preserved in result
    assert "2017" not in result and ", 2024" not in result
    assert format_resume(result, "Candidate\nEngineer at Original Corp") == result
    path = reference_template.build(result, tmp_path / "resume.docx")
    doc = Document(path)
    exported = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    exported += "\n" + "\n".join(cell.text for table in doc.tables
                                     for row in table.rows for cell in row.cells)
    assert "Added Corp #Confidential" in exported
    assert "Improved throughput by 35%" in exported
    assert "Delivered APIs with Python" in exported
    assert "2018" not in exported and "2021" not in exported
    assert "Period:" not in exported and "Dates:" not in exported
    for style in ("Normal", "ResumeName", "Headline", "Body", "Bullet", "Subheading"):
        assert doc.styles[style].font.name == "Times New Roman"


def test_inline_dates_keep_accomplishments_and_headings():
    result = format_resume("## Professional Experience\n### Engineer | Added Corp | 2020\u2013Present\n"
                           "- Launched platform in 2023 with 40% savings.\n"
                           "## Projects\n### Platform (March 2022 - June 2023)\n"
                           "- Shipped on 12/03/2023, cutting costs by 20%.")
    assert "Added Corp #Confidential" in result
    assert "2020" not in result and "2022" not in result and "2023" not in result
    assert "40% savings" in result and "20%" in result and "### Platform" in result


def test_publication_year_is_removed_without_dropping_publication():
    result = format_resume("## Publications\n- Presented at PyData Chennai, 2023.\n"
                           "## Achievements\n- Awarded in 2022.")
    assert "Presented at PyData Chennai." in result
    assert "Awarded." in result
    assert "2023" not in result and "2022" not in result


def test_confidential_markers_are_normalized_once_and_source_names_stay_public():
    resume = """## Professional Experience
### Engineer | Source Corp # confidential #Confidential
- Built services.
### Engineer | Added Corp #CONFIDENTIAL
- Deployed services.
### Engineer | New Corp
- Tested services.
## Selected Projects
### Platform
Employer: Added Corp # confidential #Confidential
Client: Source Corp
"""
    result = format_resume(resume, "Source Corp", [])
    assert result.count("Source Corp #Confidential") == 0
    assert result.count("Added Corp #Confidential") == 2
    assert result.count("New Corp #Confidential") == 1
    assert "Source Corp\n" in result
    assert "# confidential" not in result.casefold()


def test_reversed_confidential_source_role_duplicate_is_removed():
    source = "PROFESSIONAL EXPERIENCE\nSenior AI Engineer | #confidential\nWhisper platform\n"
    resume = """## Professional Experience
### Senior AI Engineer | #Confidential
- Owned model delivery and production support.
### #Confidential Whisper platform | Senior AI Engineer
- Copied role details.
### Senior AI Engineer | Added Product Division #Confidential
- Built an additional, distinct system and validated its rollout.
## Education
B.Tech
"""
    result = format_resume(resume, source)
    assert "### Senior AI Engineer | #Confidential" in result
    assert "### Senior AI Engineer | Added Product Division #Confidential" in result
    assert "### #Confidential Whisper platform | Senior AI Engineer" not in result
    assert "Copied role details" not in result
