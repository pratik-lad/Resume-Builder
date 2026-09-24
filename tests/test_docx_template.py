"""Regression checks for the deterministic, ATS-safe DOCX template."""
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "api"))

from docgen.docx_builder import render_structured  # noqa: E402
from parsers.docx_parser import parse_docx  # noqa: E402


def test_structured_template_is_parseable_single_column_and_conditional(tmp_path):
    path = tmp_path / "resume.docx"
    render_structured({
        "header": {
            "name": "Avery Patel", "location": "Pune, Maharashtra",
            "phone": "+91 90000 00000", "email": "avery@example.com",
            "linkedin": "linkedin.com/in/averypatel", "github": "github.com/averypatel",
        },
        "summary": "Early-career developer with verified Python and API project experience.",
        "skills": {"Languages": ["Python", "SQL"], "Frameworks & Tools": ["FastAPI", "Git"]},
        "experience": [{
            "title": "Software Intern", "company": "Example Labs", "location": "Pune",
            "start_date": "Jun 2024", "end_date": "Dec 2024",
            "bullets": ["Built a verified FastAPI service for internal users."],
        }],
        "projects": [{"name": "Analytics API", "technologies": "Python, FastAPI", "bullets": ["Implemented a documented REST API."]}],
        "education": [{"degree": "B.Tech Computer Science", "university": "Example University", "dates": "2021 – 2025", "cgpa": "CGPA: 8.2"}],
        # Empty sections must not create headings.
        "certifications": [], "additional": [],
    }, path)

    doc = Document(path)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "Avery Patel" in text
    assert "avery@example.com" in text
    assert "PROFESSIONAL SUMMARY" in text
    assert "TECHNICAL SKILLS" in text
    assert "CERTIFICATIONS" not in text
    assert "ADDITIONAL INFORMATION" not in text
    assert not doc.tables
    assert len(doc.sections) == 1
    assert {"ResumeName", "ContactInfo", "SectionHeading", "Subheading", "Body", "Bullet", "SkillLine", "DateLine"}.issubset({s.name for s in doc.styles})
    assert any(p.style.name == "Bullet" for p in doc.paragraphs)

    # Re-open through the project's own DOCX parser: all content stays selectable
    # and readable without relying on tables, text boxes, or images.
    parsed = parse_docx(str(path))
    assert "Avery Patel" in parsed.raw_text
    assert "Built a verified FastAPI service" in parsed.raw_text
