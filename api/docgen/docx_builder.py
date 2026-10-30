"""Deterministic, single-column DOCX resume renderer.

The optimiser supplies tailored content. This module owns every layout
decision: section order, styles, spacing, bullets, and conditional rendering.
It intentionally creates no tables, text boxes, headers, footers, or other
layout objects that make resumes harder for applicant-tracking systems to read.
"""
from __future__ import annotations

from collections import OrderedDict
from pathlib import Path
import re
from typing import Any, Iterable, Mapping

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from schemas import Fact, OptimizationStrategy, Rewrite


DEFAULT_SECTION_ORDER = (
    "summary", "skills", "experience", "projects", "education",
    "certifications", "additional",
)
SECTION_TITLES = {
    "summary": "PROFESSIONAL SUMMARY",
    "skills": "TECHNICAL SKILLS",
    "experience": "PROFESSIONAL EXPERIENCE",
    "projects": "PROJECTS",
    "education": "EDUCATION",
    "certifications": "CERTIFICATIONS",
    "additional": "ADDITIONAL INFORMATION",
}
_TYPE_TO_SECTION = {
    "summary": "summary", "experience_bullet": "experience",
    "project_bullet": "projects", "skill": "skills", "education": "education",
    "certification": "certifications", "other": "additional",
}

# Classification is deliberately small and deterministic. Unknown skills remain
# visible under "Technical Skills" rather than being guessed into a category.
_SKILL_CATEGORIES: OrderedDict[str, set[str]] = OrderedDict([
    ("Languages", {"python", "java", "javascript", "typescript", "c", "c++", "c#", "sql", "r", "go", "kotlin", "scala"}),
    ("Data & Analytics", {"excel", "power bi", "tableau", "pandas", "numpy", "data cleaning", "data visualization", "statistics", "spark", "hadoop"}),
    ("Machine Learning & AI", {"machine learning", "deep learning", "nlp", "llms", "llm", "tensorflow", "pytorch", "scikit-learn", "computer vision"}),
    ("Frameworks & Tools", {"fastapi", "django", "flask", "spring", "spring boot", "react", "node.js", "git", "docker", "kubernetes", "jupyter", "linux"}),
    ("Cloud", {"aws", "azure", "gcp", "google cloud", "amazon web services"}),
    ("Databases", {"mysql", "postgresql", "mongodb", "sqlite", "oracle", "redis", "dynamodb"}),
])


def _set_font(style: Any, size: float, bold: bool = False) -> None:
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = bold


def _set_bottom_border(paragraph: Any, color: str = "808080") -> None:
    """Add a restrained bottom border to a section heading paragraph."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def _configure_document(doc: Document) -> None:
    """Set page geometry and explicit reusable paragraph styles."""
    section = doc.sections[0]
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    # Important content is always in the body, not header/footer.
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    _set_font(normal, 9.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    definitions = {
        "ResumeName": (18.5, True, 0, 1),
        "ContactInfo": (9.0, False, 0, 5),
        "SectionHeading": (10.5, True, 9, 3),
        "Subheading": (10.0, True, 4, 0),
        "Body": (9.5, False, 0, 1),
        "Bullet": (9.5, False, 0, 1),
        "SkillLine": (9.5, False, 0, 1),
        "DateLine": (9.0, False, 0, 1),
    }
    for name, (size, bold, before, after) in definitions.items():
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        _set_font(style, size, bold)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0

    doc.styles["ResumeName"].paragraph_format.keep_with_next = True
    doc.styles["ContactInfo"].paragraph_format.keep_with_next = True
    doc.styles["SectionHeading"].paragraph_format.keep_with_next = True
    doc.styles["Subheading"].paragraph_format.keep_with_next = True
    # Inherit Word's built-in bullet numbering while exposing the required,
    # named Bullet style to callers and tests.
    doc.styles["Bullet"].base_style = doc.styles["List Bullet"]
    doc.styles["Bullet"].paragraph_format.left_indent = Inches(0.18)
    doc.styles["Bullet"].paragraph_format.first_line_indent = Inches(-0.14)


def _text(value: Any) -> str:
    return str(value or "").strip()


def _unique(items: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for item in items:
        item = _text(item)
        key = item.casefold()
        if item and key not in seen:
            seen.add(key)
            result.append(item)
    return result


def _add_heading(doc: Document, title: str) -> None:
    p = doc.add_paragraph(style="SectionHeading")
    p.add_run(title)
    _set_bottom_border(p)


def _add_bullet(doc: Document, text: str) -> None:
    if _text(text):
        doc.add_paragraph(_text(text), style="Bullet")


def _render_header(doc: Document, header: Mapping[str, Any]) -> None:
    """Render name and compact contact details in ordinary body paragraphs."""
    name = _text(header.get("name") or header.get("full_name") or "[FULL NAME]")
    p = doc.add_paragraph(style="ResumeName")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(name)

    location = _text(header.get("location"))
    if not location:
        location = ", ".join(_unique([header.get("city", ""), header.get("state", "")]))
    # `source_contacts` comes from literal source spans. Prefer it whenever
    # available so names, phone/email and profile URLs are preserved exactly.
    contacts = _unique(header.get("source_contacts", []) or [
        location, header.get("phone", ""), header.get("email", ""),
        header.get("linkedin", ""), header.get("github", "") or header.get("portfolio", ""),
    ])
    if contacts:
        p = doc.add_paragraph(" | ".join(contacts), style="ContactInfo")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _normalise_skill_groups(skills: Any) -> OrderedDict[str, list[str]]:
    """Turn mappings or fact-derived lists into non-empty labelled skill lines."""
    grouped: OrderedDict[str, list[str]] = OrderedDict()
    if isinstance(skills, Mapping):
        for label, values in skills.items():
            values = values if isinstance(values, (list, tuple, set)) else re.split(r"\s*,\s*", _text(values))
            cleaned = _unique(str(v) for v in values)
            if cleaned:
                grouped[_text(label) or "Technical Skills"] = cleaned
        return grouped

    uncategorised: list[str] = []
    for raw in skills or []:
        raw = _text(raw)
        if not raw:
            continue
        label, separator, rest = raw.partition(":")
        known_labels = {key.casefold() for key in _SKILL_CATEGORIES} | {"technical skills", "skills"}
        if separator and label.strip().casefold() in known_labels:
            grouped.setdefault(label.strip(), []).extend(_unique(re.split(r"\s*,\s*", rest)))
            continue
        for skill in _unique(re.split(r"\s*,\s*", raw)):
            category = next((name for name, terms in _SKILL_CATEGORIES.items() if skill.casefold() in terms), None)
            if category:
                grouped.setdefault(category, []).append(skill)
            else:
                uncategorised.append(skill)
    if uncategorised:
        grouped.setdefault("Technical Skills", []).extend(uncategorised)
    return OrderedDict((label, _unique(values)) for label, values in grouped.items() if _unique(values))


def _render_skills(doc: Document, skills: Any, title: str = SECTION_TITLES["skills"]) -> None:
    groups = _normalise_skill_groups(skills)
    if not groups:
        return
    _add_heading(doc, title)
    for label, values in groups.items():
        p = doc.add_paragraph(style="SkillLine")
        p.add_run(f"{label}: ").bold = True
        p.add_run(", ".join(values))


def _entry_title(entry: Mapping[str, Any], keys: tuple[str, ...]) -> str:
    return " | ".join(_unique(entry.get(key, "") for key in keys))


def _render_entries(doc: Document, section: str, entries: list[Mapping[str, Any]],
                    title: str | None = None) -> None:
    if not entries:
        return
    _add_heading(doc, title or SECTION_TITLES[section])
    for entry in entries:
        if section == "experience":
            title = _entry_title(entry, ("title", "company", "location"))
        elif section == "projects":
            title = _entry_title(entry, ("name", "technologies"))
        elif section == "education":
            title = _entry_title(entry, ("degree", "university"))
        elif section == "certifications":
            title = _entry_title(entry, ("name", "organization", "year"))
        else:
            title = _text(entry.get("title") or entry.get("text"))
        if title:
            doc.add_paragraph(title, style="Subheading")

        dates = _text(entry.get("dates"))
        if not dates:
            start, end = _text(entry.get("start_date")), _text(entry.get("end_date"))
            dates = " – ".join(part for part in (start, end or ("Present" if start else "")) if part)
        location = _text(entry.get("location"))
        if dates or (section == "education" and location):
            doc.add_paragraph(" | ".join(_unique([location, dates])), style="DateLine")

        detail = _text(entry.get("gpa") or entry.get("cgpa") or entry.get("percentage"))
        if detail:
            doc.add_paragraph(detail, style="Body")
        bullets = entry.get("bullets", [])
        if isinstance(bullets, str):
            bullets = [bullets]
        for bullet in bullets:
            _add_bullet(doc, _text(bullet))


def _render_summary(doc: Document, summary: Any,
                    title: str = SECTION_TITLES["summary"]) -> None:
    text = _text(summary)
    if text:
        _add_heading(doc, title)
        doc.add_paragraph(text, style="Body")


def _render_additional(doc: Document, items: Any,
                       title: str = SECTION_TITLES["additional"]) -> None:
    values = _unique(items if isinstance(items, (list, tuple, set)) else [_text(items)])
    if not values:
        return
    _add_heading(doc, title)
    for item in values:
        _add_bullet(doc, item)


def render_structured(data: Mapping[str, Any], out_path: Path,
                      section_order: Iterable[str] = DEFAULT_SECTION_ORDER) -> Path:
    """Render structured resume data into a stable DOCX layout.

    Empty sections are omitted. Callers may reorder valid sections for a job
    type, but cannot alter layout or styles through this function.
    """
    doc = Document()
    _configure_document(doc)
    _render_header(doc, data.get("header", {}))
    titles = {**SECTION_TITLES, **data.get("section_titles", {})}
    for section in section_order:
        if section == "summary":
            _render_summary(doc, data.get("summary", ""), titles["summary"])
        elif section == "skills":
            _render_skills(doc, data.get("skills", {}), titles["skills"])
        elif section in {"experience", "projects", "education", "certifications"}:
            entries = data.get(section, [])
            if isinstance(entries, Mapping):
                entries = [entries]
            _render_entries(doc, section, list(entries or []), titles[section])
        elif section == "additional":
            _render_additional(doc, data.get("additional", []), titles["additional"])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def _facts_to_structured(facts: list[Fact], rewrites: list[Rewrite], header_lines: list[str],
                         strategy: OptimizationStrategy | None = None) -> dict[str, Any]:
    """Adapt the existing fact ledger without changing pipeline contracts."""
    final_text = {rewrite.fact_id: rewrite.final for rewrite in rewrites}
    grouped: dict[str, list[Fact]] = {}
    order = {fid: i for i, fid in enumerate(strategy.fact_order)} if strategy else {}
    ordered_facts = sorted(enumerate(facts), key=lambda item: (order.get(item[1].fact_id, len(order) + item[0])))
    for _, fact in ordered_facts:
        grouped.setdefault(_TYPE_TO_SECTION.get(fact.type, "additional"), []).append(fact)

    header = {"name": header_lines[0] if header_lines else "[FULL NAME]"}
    if len(header_lines) > 1:
        header["source_contacts"] = header_lines[1:]

    def entries(section: str, title_key: str) -> list[dict[str, Any]]:
        by_context: OrderedDict[str, list[str]] = OrderedDict()
        for fact in grouped.get(section, []):
            by_context.setdefault(fact.context or "", []).append(final_text.get(fact.fact_id, fact.text))
        return [{title_key: context, "bullets": texts} if context else {"bullets": texts}
                for context, texts in by_context.items()]

    education = [{"degree": final_text.get(f.fact_id, f.text), "university": f.context}
                 for f in grouped.get("education", [])]
    certifications = [{"name": final_text.get(f.fact_id, f.text), "organization": f.context}
                      for f in grouped.get("certifications", [])]
    section_titles = {
        section: next((fact.section_heading for fact in section_facts if fact.section_heading), "")
        for section, section_facts in grouped.items()
    }
    return {
        "header": header,
        "section_titles": {key: value for key, value in section_titles.items() if value},
        "summary": " ".join(_unique(final_text.get(f.fact_id, f.text) for f in grouped.get("summary", []))),
        "skills": [final_text.get(f.fact_id, f.text) for f in grouped.get("skills", [])],
        "experience": entries("experience", "title"),
        "projects": entries("projects", "name"),
        "education": education,
        "certifications": certifications,
        "additional": [final_text.get(f.fact_id, f.text) for f in grouped.get("additional", [])],
    }


def build(facts: list[Fact], rewrites: list[Rewrite], header_lines: list[str], out_path: Path,
          strategy: OptimizationStrategy | None = None) -> Path:
    """Backward-compatible entry point used by the existing render stage."""
    section_order = strategy.section_order if strategy and strategy.section_order else DEFAULT_SECTION_ORDER
    return render_structured(_facts_to_structured(facts, rewrites, header_lines, strategy),
                             out_path, section_order)


def build_from_source(raw_text: str, rewrites: list[Rewrite], out_path: Path) -> Path:
    """Preserve the complete source; extracted facts are not a display inventory.

    Join PDF soft wraps within each section for exact rewrite matching. Content
    with no matching rewrite remains present, including titles and metadata.
    """
    from parsers import segmenter

    sections: list[tuple[str, list[str]]] = []
    title = ""
    lines: list[str] = []
    for line in raw_text.splitlines():
        line = line.strip()
        if not line:
            continue
        if (segmenter._LOOKUP.get(segmenter._normalise(line))
                and len(line.split()) <= 5):
            sections.append((title, lines))
            title, lines = line, []
        else:
            lines.append(line)
    sections.append((title, lines))

    doc = Document()
    _configure_document(doc)
    for title, lines in sections:
        if not title:
            if lines:
                _render_header(doc, {"name": lines[0], "source_contacts": lines[1:]})
            continue
        _add_heading(doc, title)
        canonical = segmenter._LOOKUP.get(segmenter._normalise(title))
        if canonical == "summary":
            paragraphs = [" ".join(lines)]
        else:
            paragraphs = []
            for line in lines:
                if re.fullmatch(r"[\u2022*\-]+", line):
                    continue
                # Lowercase continuations are PDF soft wraps, not new entries.
                if (paragraphs and line[0].islower()
                        and not re.match(r"^[\u2022\u2013*\-]", line)):
                    paragraphs[-1] += " " + line
                else:
                    paragraphs.append(line)
        for text in paragraphs:
            for rewrite in rewrites:
                if (canonical not in {"summary", "education", "certifications"}
                        and rewrite.status == "accepted" and rewrite.original):
                    text = text.replace(rewrite.original, rewrite.final, 1)
            if re.match(r"^[\u2022\u2013*\-]\s+", text):
                _add_bullet(doc, re.sub(r"^[\u2022\u2013*\-]\s+", "", text))
            else:
                doc.add_paragraph(text, style="Body")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def create_master_template(out_path: Path) -> Path:
    """Create a reusable, populated placeholder DOCX master template."""
    placeholder = {
        "header": {"name": "[FULL NAME]", "location": "[City, State]", "phone": "[Phone]", "email": "[Email]", "linkedin": "[LinkedIn]", "github": "[GitHub / Portfolio]"},
        "summary": "[Concise, tailored professional summary.]",
        "skills": {"Languages": ["[Language]"], "Frameworks & Tools": ["[Tool]"], "Databases": ["[Database]" ]},
        "experience": [{"title": "[JOB TITLE]", "company": "[COMPANY]", "location": "[LOCATION]", "start_date": "[START DATE]", "end_date": "[END DATE / PRESENT]", "bullets": ["[Achievement or responsibility]"]}],
        "projects": [{"name": "[PROJECT NAME]", "technologies": "[TECHNOLOGIES]", "bullets": ["[What was built and technical contribution]"]}],
        "education": [{"degree": "[DEGREE]", "university": "[UNIVERSITY]", "location": "[LOCATION]", "dates": "[DATES]"}],
        "certifications": [{"name": "[CERTIFICATION]", "organization": "[ORGANIZATION]", "year": "[YEAR]"}],
        "additional": ["[Optional award, coursework, language, or achievement]"],
    }
    return render_structured(placeholder, out_path)
