"""Presentation-only renderer inspired by the supplied Alex Example resume.

Consume existing Markdown without generating or changing candidate facts.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from docgen.docx_builder import _configure_document


SECTION_NAMES = {
    "summary", "professional summary", "profile", "professional profile",
    "objective", "career objective", "skills", "technical skills", "core skills",
    "core competencies", "education", "educational qualification", "academic qualifications",
    "experience", "work experience", "relevant experience",
    "professional experience", "employment history", "projects", "key projects",
    "selected projects", "relevant project experience", "project experience", "certifications",
    "certifications and training", "achievements", "awards", "internship",
    "internships", "additional information", "publications", "languages", "soft skills", "hobbies",
    "extra-curricular activities and opensource contribution",
}
PROJECT_FIELDS = {"title", "client", "operating system", "tools", "technologies",
                  "tools/technologies", "tools / technologies", "team size", "period",
                  "duration", "project description"}
REPAIR_HEADINGS = {"content expansion guidelines", "fabrication guidelines",
                   "quality check (perform silently before returning)", "output format"}


def plain(text):
    return text.replace("**", "").replace("__", "").strip()


def rich(paragraph, text):
    """Keep bold emphasis; never drop the text of an inline Markdown link."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    for part in re.split(r"(\*\*.*?\*\*|__.*?__)", text):
        if not part:
            continue
        bold = part.startswith(("**", "__")) and part.endswith(("**", "__"))
        run = paragraph.add_run(part[2:-2] if bold else part)
        if bold:
            run.bold = True


def configure(doc):
    _configure_document(doc)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(.52)
    section.bottom_margin = Inches(.52)
    section.left_margin = section.right_margin = Inches(.6)
    section.footer_distance = Inches(.3)
    if "Headline" not in doc.styles:
        doc.styles.add_style("Headline", WD_STYLE_TYPE.PARAGRAPH)
    for name in ("Normal", "ResumeName", "Headline", "ContactInfo", "SectionHeading",
                 "Subheading", "Body", "Bullet", "SkillLine", "DateLine"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string("263445")
        style.paragraph_format.line_spacing = 1.05
        style.paragraph_format.widow_control = True
    doc.styles["ResumeName"].font.size = Pt(25)
    doc.styles["ResumeName"].font.color.rgb = RGBColor.from_string("17324D")
    doc.styles["ResumeName"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["ResumeName"].paragraph_format.space_after = Pt(3)
    doc.styles["Headline"].font.size = Pt(11)
    doc.styles["Headline"].font.bold = True
    doc.styles["Headline"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Headline"].paragraph_format.space_after = Pt(3)
    doc.styles["Headline"].paragraph_format.keep_with_next = True
    doc.styles["ContactInfo"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["ContactInfo"].paragraph_format.keep_with_next = True
    doc.styles["SectionHeading"].font.size = Pt(10)
    doc.styles["SectionHeading"].font.color.rgb = RGBColor.from_string("246A7B")
    doc.styles["Subheading"].font.size = Pt(11)
    doc.styles["Subheading"].font.underline = False
    doc.styles["Subheading"].paragraph_format.space_before = Pt(8)
    doc.styles["Subheading"].paragraph_format.space_after = Pt(4)
    doc.styles["Subheading"].paragraph_format.keep_with_next = True
    doc.styles["Body"].paragraph_format.space_after = Pt(2)
    doc.styles["Bullet"].paragraph_format.space_after = Pt(2)
    doc.styles["DateLine"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.styles["DateLine"].paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def section_heading(doc, title):
    paragraph = doc.add_paragraph(style="SectionHeading")
    rich(paragraph, title.upper())
    borders = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    for key, value in {"val": "single", "sz": "4", "space": "5", "color": "CBD8DF"}.items():
        border.set(qn(f"w:{key}"), value)
    borders.append(border)
    paragraph._p.get_or_add_pPr().append(borders)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(6)


def table(doc, rows, header=False):
    columns = max(len(row) for row in rows)
    result = doc.add_table(rows=0, cols=columns)
    result.style = "Normal Table"
    result.autofit = False
    widths = [2.1, 5.2] if columns == 2 else [7.3 / columns] * columns
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    for index, values in enumerate(rows):
        row = result.add_row()
        # Keep each row together; long tables can still flow between pages.
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if index == 0 and header:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for col, cell in enumerate(row.cells):
            cell.width = Inches(widths[col])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.space_before = Pt(4)
            if header and index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EFF4F7")
                cell._tc.get_or_add_tcPr().append(shade)
            rich(p, values[col] if col < len(values) else "")
            if (header and index == 0) or (not header and col == 0):
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def role_date_row(doc, role, dates):
    """Align role and tenure like the reference without using floating text."""
    row = doc.add_table(rows=1, cols=2)
    row.autofit = False
    row.columns[0].width, row.columns[1].width = Inches(5.5), Inches(1.8)
    row.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for cell, width, value in zip(row.rows[0].cells, (5.5, 1.8), (role, dates)):
        cell.width = Inches(width)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if value == dates else
                               WD_ALIGN_PARAGRAPH.LEFT)
        rich(paragraph, value)
        for run in paragraph.runs:
            run.bold = True


def group_projects_by_employer(markdown: str) -> str:
    """DOCX-only layout transform; leave the generated Markdown and claims intact."""
    experience = re.search(r"(?ms)^## Professional Experience\s*\n(.*?)(?=^## |\Z)", markdown)
    projects = re.search(r"(?ms)^## Selected Projects\s*\n(.*?)(?=^## |\Z)", markdown)
    if not experience or not projects:
        return markdown
    role_parts = re.split(r"(?m)(?=^### )", experience.group(1))
    project_parts = re.split(r"(?m)(?=^### )", projects.group(1))
    grouped = {index: [] for index in range(len(role_parts))}
    remaining = []
    for project in project_parts:
        employer = re.search(r"(?im)^Employer:\s*(.+)$", project)
        if not employer:
            if project.strip():
                remaining.append(project)
            continue
        matching = next((index for index, role in enumerate(role_parts)
                         if re.search(r"(?m)^### [^\n]*\|\s*" +
                                      re.escape(employer.group(1).strip()) + r"\s*$", role)), None)
        if matching is None:
            remaining.append(project)
            continue
        grouped[matching].append(re.sub(r"(?im)^Employer:\s*.+\n?", "", project, count=1).strip())
    for index, attached in grouped.items():
        if attached:
            role_parts[index] = role_parts[index].rstrip() + "\n\n" + "\n\n".join(attached) + "\n\n"
    revised_experience = "".join(role_parts).strip()
    revised_projects = "\n\n".join(remaining).strip()
    result = markdown[:experience.start(1)] + revised_experience + "\n" + markdown[experience.end(1):]
    # Locate again because the preceding section changed length.
    projects = re.search(r"(?ms)^## Selected Projects\s*\n(.*?)(?=^## |\Z)", result)
    if projects:
        result = (result[:projects.start()] +
                  ("## Selected Projects\n" + revised_projects + "\n\n" if revised_projects else "") +
                  result[projects.end():])
    return result


def build(markdown: str, out_path: Path) -> Path:
    doc = Document()
    configure(doc)
    lines = group_projects_by_employer(markdown).splitlines()
    current_section = ""
    has_name = False
    skip_repair_text = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line or line.startswith("```") or re.fullmatch(r"[-*_]{3,}", line):
            continue
        # Preserve Markdown tables as actual DOCX tables.
        if "|" in line and i < len(lines) and re.fullmatch(r"[\s|:\-]+", lines[i].strip()) and "-" in lines[i]:
            rows = [[cell.strip() for cell in line.strip("|").split("|")]]
            i += 1
            while i < len(lines) and "|" in lines[i]:
                rows.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
                i += 1
            table(doc, rows, header=True)
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", line)
        text = heading.group(2) if heading else line
        key = plain(text).rstrip(":").casefold()
        if heading and key in REPAIR_HEADINGS:
            skip_repair_text = True
            continue
        if skip_repair_text:
            if not (heading and len(heading.group(1)) <= 2 and key in SECTION_NAMES):
                continue
            skip_repair_text = False
        if key in SECTION_NAMES or (heading and len(heading.group(1)) == 2):
            current_section = key
            section_heading(doc, plain(text))
            continue
        if heading and len(heading.group(1)) == 3 and current_section == "professional experience":
            next_line = lines[i].strip() if i < len(lines) else ""
            if re.match(r"(?i)^Dates:\s*\S", next_line):
                role_date_row(doc, plain(text), next_line.split(":", 1)[1].strip())
                i += 1
                continue
        if not has_name and not current_section:
            rich(doc.add_paragraph(style="ResumeName"), text)
            has_name = True
            continue
        if not current_section:
            is_role = bool(re.search(r"\b(?:engineer|developer|analyst|manager|designer|scientist|architect)\b",
                                     text, re.I)) and not re.search(r"[@|]|https?://|\bwww\.", text, re.I)
            is_contact = bool(re.search(
                r"@|https?://|\b(?:linkedin|github)\b|(?:^|\s)\+?\d[\d\s().-]{7,}\d(?:\s|$)",
                text, re.I))
            # Only identity/contact lines belong above the first section. Do not
            # promote source summary or education spillover into the header.
            if has_name and not is_role and not is_contact:
                continue
            p = doc.add_paragraph(style="Headline" if is_role else "ContactInfo")
            rich(p, text)
            if not is_role:
                for run in p.runs:
                    run.font.color.rgb = RGBColor.from_string("536779")
            continue
        # Skills are compact single-line categories; only explicit Markdown
        # tables elsewhere remain tables in the exported document.
        cleaned = plain(re.sub(r"^(?:[-*+•]|â€¢)\s+", "", line))
        pair = cleaned.split(":", 1)
        is_skills = "skill" in current_section or current_section == "core competencies"
        if len(pair) == 2 and (is_skills or pair[0].casefold() in PROJECT_FIELDS):
            rows = [pair]
            while i < len(lines):
                next_line = plain(re.sub(r"^(?:[-*+•]|â€¢)\s+", "", lines[i].strip()))
                next_pair = next_line.split(":", 1)
                if len(next_pair) != 2 or not (is_skills or next_pair[0].casefold() in PROJECT_FIELDS):
                    break
                rows.append(next_pair)
                i += 1
            for label, value in rows:
                p = doc.add_paragraph(style="SkillLine" if is_skills else "Body")
                run = p.add_run(label.strip() + ": ")
                run.bold = True
                run.font.color.rgb = RGBColor.from_string("17324D")
                rich(p, value.strip())
            continue
        if current_section == "education" and not heading:
            rows = []
            while True:
                rows.append([re.sub(r"^(?:[-*+•]|â€¢)\s+", "", line)])
                if i >= len(lines) or not lines[i].strip() or lines[i].lstrip().startswith("#") or plain(lines[i]).rstrip(":").casefold() in SECTION_NAMES:
                    break
                line = lines[i].strip()
                i += 1
            for row in rows:
                rich(doc.add_paragraph(style="Body"), row[0])
            continue
        bullet = re.match(r"^(?:[-*+•]\s+|â€¢\s+|\d+[.)]\s+)(.*)", line)
        if heading:
            rich(doc.add_paragraph(style="Subheading"), text)
        elif current_section in {"projects", "selected projects", "key projects", "project experience",
                                 "relevant project experience"} and key in {
                                     "role and responsibilities", "technical", "functional"}:
            paragraph = doc.add_paragraph(style="Body")
            paragraph.add_run(plain(text).rstrip(":") + ":").bold = True
        elif re.match(r"^(?:Dates?|Period|Duration):\s*\S", line, re.I):
            rich(doc.add_paragraph(style="DateLine"), line)
        elif bullet:
            rich(doc.add_paragraph(style="Bullet"), bullet.group(1))
        elif current_section in {"projects", "selected projects", "key projects", "project experience",
                                 "relevant project experience"} and i < len(lines) \
                and (i >= 2 and not lines[i - 2].strip()
                     or doc.paragraphs[-1].style.name == "SectionHeading") \
                and len(line) <= 90 and not line.endswith((".", ":")):
            rich(doc.add_paragraph(style="Subheading"), line)
        elif line.startswith("**") and line.endswith("**"):
            rich(doc.add_paragraph(style="Subheading"), line)
        else:
            rich(doc.add_paragraph(style="Body"), line)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
