"""DOCX extraction that preserves paragraph/table document order."""
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from parsers.pdf_parser import Line, ParsedDoc


def parse_docx(path: str) -> ParsedDoc:
    doc = Document(path)
    lines: list[Line] = []
    cursor = 0

    def append_paragraph(para):
        nonlocal cursor
        text = para.text.strip()
        if not text:
            return
        style = ((para.style.name if para.style else "") or "").lower()
        bold = "heading" in style or any(r.bold for r in para.runs if r.bold is not None)
        size = 14.0 if "heading" in style else 10.0
        ln = Line(text=text, page=1, x0=0.0, y0=float(cursor), size=size, bold=bool(bold))
        ln.char_start = cursor
        ln.char_end = cursor + len(text)
        cursor = ln.char_end + 1
        lines.append(ln)

    def append_table(table):
        nonlocal cursor
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if not cells:
                continue
            text = " | ".join(dict.fromkeys(cells))
            ln = Line(text=text, page=1, x0=0.0, y0=float(cursor), size=10.0, bold=False)
            ln.char_start = cursor
            ln.char_end = cursor + len(text)
            cursor = ln.char_end + 1
            lines.append(ln)

    # ``doc.paragraphs`` followed by ``doc.tables`` loses chronology when a
    # resume alternates narrative sections and skill/education tables. Walk
    # the body XML so headings, paragraphs, and tables stay in source order.
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            append_paragraph(Paragraph(child, doc))
        elif child.tag.endswith("}tbl"):
            append_table(Table(child, doc))

    raw = "\n".join(l.text for l in lines)
    return ParsedDoc(lines=lines, raw_text=raw, body_size=10.0)
